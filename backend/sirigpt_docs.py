"""
Document helpers for the V3 Document Engine.

Uploaded into the Code Interpreter sandbox (the model builds and edits Word /
Excel / PowerPoint files with it) AND imported by the server (doc_engine.py)
to check every file the sandbox returns. One reviewed module, so Thai fonts,
page numbers, slide layout and the file checks behave the same on every
request instead of being re-invented by the model.

The sandbox stores uploads as /mnt/data/file-<id>-<name> (seen on the real
container), so the module is loaded by pattern and user_files() gives the
attached files back under their own names.

Sandbox use (doc_engine.LOADER):
    import glob, shutil, sys; shutil.copy(sorted(glob.glob("/mnt/data/*sirigpt_docs.py"))[0], "/tmp/sirigpt_docs.py")
    sys.path.insert(0, "/tmp"); import sirigpt_docs as sd
    doc = sd.new_docx()                       # A4, Thai-ready styles, page numbers
    sd.heading(doc, "รายงานยอดขาย", 1); sd.para(doc, "…")
    sd.table(doc, [["เดือน", "ยอด"], ["ม.ค.", 41650]]); sd.picture(doc, "/mnt/data/chart.png")
    out = sd.output_path("รายงาน.docx"); doc.save(out); print(sd.check(out))

python-docx / python-pptx / openpyxl are imported lazily, so importing this
module never fails; only the helpers that need a library do.
"""

from __future__ import annotations

import math
import re
import zipfile
from pathlib import Path

DOCX_FONT = "Sarabun"          # embedded into the .docx by the server, so it shows everywhere
PPTX_FONT = "Leelawadee UI"    # ships with Windows; PowerPoint can't take embedded fonts from us
THAI_FONTS = {
    "sarabun", "th sarabun new", "th sarabunpsk", "leelawadee", "leelawadee ui", "tahoma",
    "angsana new", "angsanaupc", "cordia new", "cordiaupc", "browallia new", "browalliaupc",
    "loma", "norasi", "garuda", "kinnari", "noto sans thai", "noto serif thai", "prompt", "kanit",
}
DATA_DIR = "/mnt/data"
_UPLOAD_PREFIX = re.compile(r"^file-[A-Za-z0-9]{8,}-")
_SUPPORT_NAMES = {"sirigpt_docs.py", "sirigpt_tools.py"}
_THAI = re.compile(r"[฀-๿]")
_THAI_MARKS = set("ัิีึืฺุู็่้๊๋์ํ๎")


# ════════════════════════════════════════════════════════════ sandbox files ══

def clean_name(name: str) -> str:
    """'file-3fWZXgXr91Vu6UHy2Je415-รายงาน.docx' → 'รายงาน.docx'."""
    return _UPLOAD_PREFIX.sub("", Path(name).name)


def user_files(folder: str = DATA_DIR) -> dict[str, str]:
    """The user's attached files as {own name: path}; our module and fonts left out.
    Call it before saving results — files saved later would be listed too."""
    out = {}
    for p in sorted(Path(folder).iterdir()):
        name = clean_name(p.name)
        if p.is_file() and name not in _SUPPORT_NAMES and not name.lower().endswith(".ttf"):
            out[name] = str(p)
    return out


def output_path(name: str, folder: str = DATA_DIR) -> str:
    """Where to save a result for the user: /mnt/data/<name>, never an upload's prefixed path."""
    return str(Path(folder) / clean_name(name))


def font_path(style: str = "Regular", folder: str = DATA_DIR) -> str | None:
    hits = sorted(Path(folder).glob(f"*Sarabun-{style}.ttf"))
    return str(hits[0]) if hits else None


def thai_matplotlib(folder: str = DATA_DIR) -> None:
    """Register the uploaded Sarabun fonts with matplotlib so Thai chart labels render."""
    import matplotlib
    from matplotlib import font_manager

    for p in sorted(Path(folder).glob("*Sarabun-*.ttf")):
        font_manager.fontManager.addfont(str(p))
    matplotlib.rcParams["font.family"] = "Sarabun"
    matplotlib.rcParams["axes.unicode_minus"] = False


# ═════════════════════════════════════════════════════════════════════ Word ══

def _w(tag: str):
    from docx.oxml.ns import qn

    return qn(tag)


# Word rejects a file whose run properties are out of schema order, so new
# elements are inserted at their place in CT_RPr's sequence.
_RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
    "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
    "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
    "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]


def _insert_ordered(rpr, el) -> None:
    rank = {_w("w:" + t): i for i, t in enumerate(_RPR_ORDER)}
    mine = rank.get(el.tag, len(_RPR_ORDER))
    for i, child in enumerate(rpr):
        if rank.get(child.tag, -1) > mine:
            rpr.insert(i, el)
            return
    rpr.append(el)


def _set_rfonts(rpr, font: str) -> None:
    """Latin, Thai (complex script) and East-Asian slots — Word renders Thai
    from the *cs* slot, which python-docx's font.name does not set."""
    from docx.oxml import OxmlElement

    rf = rpr.find(_w("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        _insert_ordered(rpr, rf)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(_w(theme), None)      # a theme font would override ours
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(_w(slot), font)


def _set_val(rpr, tag: str, val: str | None) -> None:
    from docx.oxml import OxmlElement

    el = rpr.find(_w(tag))
    if el is None:
        el = OxmlElement(tag)
        _insert_ordered(rpr, el)
    if val is not None:
        el.set(_w("w:val"), val)


def thai_font(target, font: str = DOCX_FONT, size: float | None = None, bold: bool | None = None,
              color: str | None = None) -> None:
    """Font for a style or a run, with the Thai slot and Thai size/bold set too."""
    from docx.shared import Pt, RGBColor

    rpr = target._element.get_or_add_rPr()
    _set_rfonts(rpr, font)
    if size:
        target.font.size = Pt(size)
        _set_val(rpr, "w:szCs", str(int(round(size * 2))))
    if bold is not None:
        target.font.bold = bold
        _set_val(rpr, "w:bCs", "1" if bold else "0")
    if color:
        target.font.color.rgb = RGBColor.from_string(color)
    _set_val(rpr, "w:lang", None)
    rpr.find(_w("w:lang")).set(_w("w:bidi"), "th-TH")


def new_docx(template: str | None = None, font: str = DOCX_FONT, size: float = 14,
             margins_cm: float = 2.0, page_numbers: bool = True):
    """A4 document with Thai-ready styles (Normal, headings, title, caption,
    lists, table text) and "หน้า X จาก Y" in the footer. With a template, its
    page setup and styles are kept and only the fonts are made Thai-ready."""
    from docx import Document
    from docx.shared import Cm

    doc = Document(template) if template else Document()
    if not template:
        for s in doc.sections:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(margins_cm)
    sizes = {"Normal": size, "Title": size + 12, "Heading 1": size + 6, "Heading 2": size + 3,
             "Heading 3": size + 1, "Caption": size - 2, "List Bullet": size, "List Number": size}
    for name, pt in sizes.items():
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        thai_font(style, font, None if template else pt,
                  bold=True if name.startswith(("Heading", "Title")) else None,
                  color="3B1F6B" if name.startswith(("Heading", "Title")) and not template else None)
    if page_numbers:
        add_page_numbers(doc, font=font)
    return doc


def _field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement

    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(_w("w:fldCharType"), kind)
            run._r.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            run.text = text


def add_page_numbers(doc, text: str = "หน้า {PAGE} จาก {NUMPAGES}", font: str = DOCX_FONT) -> None:
    """Centered page numbers in every section's footer (real PAGE/NUMPAGES fields)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for part in re.split(r"(\{PAGE\}|\{NUMPAGES\})", text):
            if part == "{PAGE}":
                _field(p, "PAGE")
            elif part == "{NUMPAGES}":
                _field(p, "NUMPAGES")
            elif part:
                p.add_run(part)
        for r in p.runs:
            thai_font(r, font, 11)


def heading(doc, text: str, level: int = 1):
    return doc.add_heading(_plain(text), level=level)


def para(doc, text: str = "", bold: bool = False, align: str | None = None, size: float | None = None,
         style: str | None = None):
    """Paragraph; **bold** segments in `text` are made bold."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = doc.add_paragraph(style=style)
    _md_runs(p, text, bold, size)
    if align:
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    return p


def _md_runs(p, text: str, bold: bool = False, size: float | None = None) -> None:
    """**bold** segments become bold runs (no literal asterisks)."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            r = p.add_run(_plain(part))
            if bold or i % 2:
                r.bold = True
                _set_val(r._r.get_or_add_rPr(), "w:bCs", "1")
            if size:
                thai_font(r, DOCX_FONT, size)


def bullets(doc, items: list[str], numbered: bool = False) -> None:
    """Bulleted (or numbered) list; **bold** works as in para()."""
    for it in items:
        _md_runs(doc.add_paragraph(style="List Number" if numbered else "List Bullet"), str(it))


def _fmt_cell(v) -> tuple[str, bool]:
    if isinstance(v, bool) or v is None:
        return ("" if v is None else str(v)), False
    if isinstance(v, int):
        return f"{v:,}", True
    if isinstance(v, float):
        return (f"{v:,.2f}" if not v.is_integer() else f"{int(v):,}"), True
    return _plain(str(v)), False


def _plain(text: str) -> str:
    """Drop markdown bold markers where we can't show bold (titles, table cells)."""
    return text.replace("**", "")


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement

    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(_w("w:val"), "clear")
    shd.set(_w("w:color"), "auto")
    shd.set(_w("w:fill"), fill)
    tcpr.append(shd)


def table(doc, rows: list[list], header: bool = True, col_widths_cm: list[float] | None = None,
          style: str = "Table Grid", font_size: float | None = 12):
    """Table from rows (first row = header). Numbers are right-aligned with
    thousands separators; the header is shaded, bold and repeats on each page."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm

    n_cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=n_cols)
    try:
        t.style = doc.styles[style]
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(n_cols):
            text, numeric = _fmt_cell(row[ci] if ci < len(row) else "")
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            thai_font(run, DOCX_FONT, font_size, bold=True if (header and ri == 0) else None)
            if header and ri == 0:
                _shade(cell, "E4DDF1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if col_widths_cm and ci < len(col_widths_cm):
                cell.width = Cm(col_widths_cm[ci])
        if header and ri == 0:
            trpr = t.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(_w("w:val"), "true")
            trpr.append(th)
    return t


def picture(doc, path: str, width_cm: float = 15.0, caption: str | None = None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        try:
            c = doc.add_paragraph(caption, style="Caption")
        except KeyError:
            c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def page_break(doc) -> None:
    from docx.enum.text import WD_BREAK

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ═══════════════════════════════════════════════════════════════ PowerPoint ══

EMU_PER_PT = 12700
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5          # 16:9
_ACCENT = "5B2C9F"
_DARK = "1F1235"
_PALETTE = ["5B2C9F", "E0A100", "2A9D8F", "D1495B", "4C6EF5"]


def _pptx_font(run_or_font, font: str = PPTX_FONT, size: float | None = None, bold: bool | None = None,
               color: str | None = None) -> None:
    """Latin, East-Asian and complex-script (Thai) typefaces on a run."""
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    f = run_or_font.font if hasattr(run_or_font, "font") else run_or_font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)
    rpr = f._rPr if hasattr(f, "_rPr") else f._element
    # latin, ea, cs must follow the fill elements and precede sym/hlink*/rtl/extLst
    for tag in ("a:latin", "a:ea", "a:cs"):
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
    after = next((c for c in rpr if c.tag in {qn(t) for t in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver",
                                                               "a:rtl", "a:extLst")}), None)
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.makeelement(qn(tag), {"typeface": font})
        if after is not None:
            after.addprevious(el)
        else:
            rpr.append(el)


def text_width_em(text: str) -> float:
    """Rough width of `text` in ems (Thai marks take no width)."""
    w = 0.0
    for ch in text:
        if ch in _THAI_MARKS:
            continue
        if ch == " ":
            w += 0.3
        elif _THAI.match(ch):
            w += 0.58
        elif ch.isupper() or ch.isdigit():
            w += 0.62
        else:
            w += 0.52
    return w


def estimate_lines(paragraphs: list[str], size_pt: float, width_pt: float) -> int:
    per_line = max(width_pt / size_pt, 1.0)
    return sum(max(1, math.ceil(text_width_em(p) / per_line)) for p in paragraphs)


def fit_font_size(paragraphs: list[str], width_pt: float, height_pt: float,
                  max_size: float = 28, min_size: float = 12, line_spacing: float = 1.25) -> float:
    """Largest size (step 1pt) whose estimated wrapped height fits the box."""
    size = max_size
    while size > min_size:
        lines = estimate_lines(paragraphs, size, width_pt)
        if lines * size * line_spacing + size * 0.3 * len(paragraphs) <= height_pt:
            return size
        size -= 1
    return min_size


def new_pptx(template: str | None = None):
    """16:9 deck (or the user's template, keeping its size/layouts/theme)."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation(template) if template else Presentation()
    if not template:
        prs.slide_width, prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
    return prs


def _blank(prs):
    layouts = prs.slide_layouts
    for layout in layouts:
        if layout.name.lower() == "blank":
            return prs.slides.add_slide(layout)
    return prs.slides.add_slide(layouts[min(6, len(layouts) - 1)])


def _box(slide, left_in, top_in, width_in, height_in):
    from pptx.util import Inches

    tb = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def _slide_dims(prs) -> tuple[float, float]:
    return prs.slide_width / 914400, prs.slide_height / 914400


def _title(prs, slide, text: str, size: float = 32) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    w, _ = _slide_dims(prs)
    tb, tf = _box(slide, 0.6, 0.35, w - 1.2, 1.0)
    text = _plain(text)
    size = fit_font_size([text], (w - 1.4) * 72, 0.95 * 72, max_size=size, min_size=20)
    tf.text = text
    _pptx_font(tf.paragraphs[0].runs[0], size=size, bold=True, color=_DARK)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.38), Inches(1.6), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(_ACCENT)
    bar.line.fill.background()


def title_slide(prs, title: str, subtitle: str | None = None):
    from pptx.enum.text import PP_ALIGN

    s = _blank(prs)
    w, h = _slide_dims(prs)
    title, subtitle = _plain(title), subtitle and _plain(subtitle)
    tb, tf = _box(s, 0.8, h * 0.32, w - 1.6, 1.6)
    tf.text = title
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _pptx_font(tf.paragraphs[0].runs[0], size=fit_font_size([title], (w - 1.8) * 72, 1.5 * 72, 44, 28),
               bold=True, color=_DARK)
    if subtitle:
        tb2, tf2 = _box(s, 0.8, h * 0.32 + 1.7, w - 1.6, 1.0)
        tf2.text = subtitle
        tf2.paragraphs[0].alignment = PP_ALIGN.CENTER
        _pptx_font(tf2.paragraphs[0].runs[0], size=fit_font_size([subtitle], (w - 1.8) * 72, 0.9 * 72, 24, 14),
                   color=_ACCENT)
    return s


def bullet_slide(prs, title: str, items: list, notes: str | None = None, max_size: float = 30):
    """Title + bullets; items are str or (str, level). The font shrinks to fit;
    more than ~7 long bullets should be split over two slides."""
    s = _blank(prs)
    _title(prs, s, title)
    w, h = _slide_dims(prs)
    tb, tf = _box(s, 0.8, 1.7, w - 1.6, h - 2.3)
    _bullets_into(tf, items, (w - 1.8) * 72, (h - 2.4) * 72, max_size)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def _bullets_into(tf, items: list, width_pt: float, height_pt: float, max_size: float) -> None:
    """Bullets sized to fit; **bold** segments become bold runs (no literal asterisks)."""
    from pptx.util import Pt

    texts = [_plain(i if isinstance(i, str) else i[0]) for i in items]
    size = fit_font_size(["• " + t for t in texts], width_pt, height_pt, max_size, 14)
    for n, it in enumerate(items):
        text, level = (it, 0) if isinstance(it, str) else (it[0], it[1])
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.level = level
        p.space_after = Pt(round(size * 0.3))
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", ("• " if level == 0 else "– ") + text)):
            if part:
                r = p.add_run()
                r.text = _plain(part)
                _pptx_font(r, size=size - 2 * level, bold=True if i % 2 else None, color=_DARK)


def kpi_slide(prs, title: str, items: list, points: list | None = None, notes: str | None = None):
    """Up to 4 key numbers as cards [(label, value), …], optionally with bullets
    below — instead of hand-placed boxes, which landed on top of other text."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches

    s = _blank(prs)
    _title(prs, s, title)
    w, h = _slide_dims(prs)
    items = items[:4]
    gap = 0.4
    card_w = (w - 1.2 - gap * (len(items) - 1)) / max(1, len(items))
    top, card_h = (1.8, 1.8) if points else (2.4, 2.6)
    for i, (label, value) in enumerate(items):
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + i * (card_w + gap)), Inches(top),
                                  Inches(card_w), Inches(card_h))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string("F3EFFA")
        card.line.color.rgb = RGBColor.from_string(_PALETTE[i % len(_PALETTE)])
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        value_text = _fmt_cell(value)[0]
        label_text = _plain(str(label))
        for n, (text, size, bold, color) in enumerate((
            (label_text, fit_font_size([label_text], (card_w - 0.3) * 72, card_h * 0.3 * 72, 20, 12), None, "555555"),
            (value_text, fit_font_size([value_text], (card_w - 0.3) * 72, card_h * 0.45 * 72, 40, 16), True, _ACCENT),
        )):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = text
            _pptx_font(r, size=size, bold=bold, color=color)
    if points:
        tb, tf = _box(s, 0.8, top + card_h + 0.3, w - 1.6, h - (top + card_h + 0.8))
        _bullets_into(tf, points, (w - 1.8) * 72, (h - (top + card_h + 0.9)) * 72, 24)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def table_slide(prs, title: str, rows: list[list], font_size: float | None = None):
    """Native (editable) table; font shrinks with the row count."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    s = _blank(prs)
    _title(prs, s, title)
    w, h = _slide_dims(prs)
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    avail_h = h - 2.2
    size = font_size or max(10, min(22, int(avail_h * 72 / n_rows / 1.9)))
    shape = s.shapes.add_table(n_rows, n_cols, Inches(0.6), Inches(1.7), Inches(w - 1.2),
                               Inches(min(avail_h, n_rows * size * 1.9 / 72)))
    t = shape.table
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            text, numeric = _fmt_cell(row[ci] if ci < len(row) else "")
            cell = t.cell(ri, ci)
            cell.text = text
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if numeric and ri else (PP_ALIGN.CENTER if ri == 0 else PP_ALIGN.LEFT)
            _pptx_font(p.runs[0] if p.runs else p.font, size=size, bold=ri == 0,
                       color="FFFFFF" if ri == 0 else _DARK)
            if ri == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(_ACCENT)
    return s


def image_slide(prs, title: str, image_path: str, caption: str | None = None):
    """Picture fitted inside the content area (aspect ratio kept) + caption."""
    from PIL import Image
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    s = _blank(prs)
    _title(prs, s, title)
    w, h = _slide_dims(prs)
    box_w, box_h = w - 1.6, h - (2.6 if caption else 2.1)
    with Image.open(image_path) as im:
        ratio = im.width / im.height
    pic_w = min(box_w, box_h * ratio)
    pic_h = pic_w / ratio
    s.shapes.add_picture(str(image_path), Inches((w - pic_w) / 2), Inches(1.7), Inches(pic_w), Inches(pic_h))
    if caption:
        tb, tf = _box(s, 0.8, 1.75 + pic_h, w - 1.6, 0.6)
        tf.text = caption
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _pptx_font(tf.paragraphs[0].runs[0], size=14, color="555555")
    return s


def chart_slide(prs, title: str, categories: list, series: dict, kind: str = "column",
                number_format: str = "#,##0"):
    """Native PowerPoint chart (editable in PowerPoint, data inside the file).
    kind: column | bar | line | pie."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    s = _blank(prs)
    _title(prs, s, title)
    w, h = _slide_dims(prs)
    data = CategoryChartData()
    data.categories = [str(c) for c in categories]
    for name, values in series.items():
        data.add_series(str(name), [float(v) for v in values])
    ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[kind]
    gf = s.shapes.add_chart(ctype, Inches(0.8), Inches(1.7), Inches(w - 1.6), Inches(h - 2.3), data)
    chart = gf.chart
    chart.has_legend = len(series) > 1 or kind == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        _pptx_font(chart.legend.font, size=14)
    plot = chart.plots[0]
    if kind != "pie":                                  # brand colours, one per series
        from pptx.dml.color import RGBColor

        for n, ser in enumerate(plot.series):
            fill = ser.format.line if kind == "line" else ser.format.fill
            colour = RGBColor.from_string(_PALETTE[n % len(_PALETTE)])
            if kind == "line":
                fill.color.rgb = colour
            else:
                fill.solid()
                fill.fore_color.rgb = colour
    plot.has_data_labels = True
    plot.data_labels.number_format = number_format
    plot.data_labels.number_format_is_linked = False
    _pptx_font(plot.data_labels.font, size=12)
    _pptx_font(chart.font, size=14)
    chart.font.size = Pt(14)
    return s


# ═══════════════════════════════════════════════════════════════════ Excel ══

def prepare_workbook(wb) -> None:
    """Excel / LibreOffice recalculate every formula when the file opens
    (openpyxl stores formulas without results)."""
    from openpyxl.workbook.properties import CalcProperties

    if wb.calculation is None:
        wb.calculation = CalcProperties()
    wb.calculation.fullCalcOnLoad = True


def style_table(ws, header_row: int = 1, number_format: str = "#,##0.00", font: str = DOCX_FONT) -> None:
    """Bold shaded header, number format, column widths, frozen header."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widths: dict[int, float] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            c.font = Font(name=font, bold=c.row == header_row, size=11)
            if c.row == header_row:
                c.fill = PatternFill("solid", fgColor="E4DDF1")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                c.number_format = number_format
            elif isinstance(c.value, str) and c.value.startswith("="):
                c.number_format = number_format
            widths[c.column] = max(widths.get(c.column, 0), min(50, text_width_em(str(c.value)) * 1.3 + 2))
    for col, wdt in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, wdt)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


# ── a small formula evaluator, to check the numbers a workbook will show ────

class _Unsupported(Exception):
    pass


class _XLError(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


_TOKEN = re.compile(r"""\s*(?:
    (?P<str>"(?:[^"]|"")*")|
    (?P<func>[A-Za-z][A-Za-z0-9.]*)\s*\(|
    (?P<bool>TRUE|FALSE)\b|
    (?P<ref>(?:(?:'(?:[^']|'')+'|[^\s!'"(),:;+\-*/^&=<>%]+)!)?\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?)|
    (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)|
    (?P<err>\#[A-Z0-9/!?]+)|
    (?P<op><>|<=|>=|[-+*/^&=<>(),;%])
)""", re.X)


def _tokens(formula: str) -> list[tuple[str, str]]:
    out, pos, s = [], 0, formula.strip()
    while pos < len(s):
        m = _TOKEN.match(s, pos)
        if not m or m.end() == pos:
            raise _Unsupported(f"cannot read formula near {s[pos:pos + 12]!r}")
        kind = m.lastgroup
        out.append((kind, m.group(kind)))
        pos = m.end()
    return out


def _cell_xy(ref: str) -> tuple[int, int]:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    col, row = coordinate_from_string(ref.replace("$", ""))
    return row, column_index_from_string(col)


class _Evaluator:
    def __init__(self, wb):
        self.wb = wb
        self.memo: dict[tuple[str, str], object] = {}
        self.busy: set[tuple[str, str]] = set()

    # values ----------------------------------------------------------------
    def cell(self, sheet: str, ref: str):
        key = (sheet, ref.replace("$", "").upper())
        if key in self.memo:
            return self.memo[key]
        if sheet not in self.wb.sheetnames:
            raise _XLError("#REF!")
        v = self.wb[sheet][key[1]].value
        if isinstance(v, str) and v.startswith("="):
            if key in self.busy:
                raise _Unsupported("circular reference")
            self.busy.add(key)
            try:
                v = self.formula(sheet, v[1:])
            finally:
                self.busy.discard(key)
        elif not isinstance(v, (int, float, str, bool)) and v is not None:
            raise _Unsupported(f"value type {type(v).__name__}")
        self.memo[key] = v
        return v

    def range(self, sheet: str, a: str, b: str) -> list:
        r1, c1 = _cell_xy(a)
        r2, c2 = _cell_xy(b)
        if (r2 - r1 + 1) * (c2 - c1 + 1) > 200_000:
            raise _Unsupported("range too large")
        from openpyxl.utils import get_column_letter

        return [self.cell(sheet, f"{get_column_letter(c)}{r}")
                for r in range(min(r1, r2), max(r1, r2) + 1) for c in range(min(c1, c2), max(c1, c2) + 1)]

    # parser (precedence: compare < & < +- < */ < ^ < unary < %) -----------
    def formula(self, sheet: str, text: str):
        # a referenced cell's formula is parsed in the middle of this one
        saved = (getattr(self, "sheet", None), getattr(self, "toks", None), getattr(self, "i", 0))
        self.sheet, self.toks, self.i = sheet, _tokens(text), 0
        try:
            val = self.compare()
            if self.i != len(self.toks):
                raise _Unsupported("unexpected text in formula")
            return val
        finally:
            self.sheet, self.toks, self.i = saved

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def take(self):
        t = self.peek()
        self.i += 1
        return t

    def compare(self):
        a = self.concat()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            b = self.concat()
            a = {"=": a == b, "<>": a != b, "<": _n(a) < _n(b), ">": _n(a) > _n(b),
                 "<=": _n(a) <= _n(b), ">=": _n(a) >= _n(b)}[op]
        return a

    def concat(self):
        a = self.additive()
        while self.peek()[1] == "&":
            self.take()
            a = _s(a) + _s(self.additive())
        return a

    def additive(self):
        a = self.mult()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            b = self.mult()
            a = _n(a) + _n(b) if op == "+" else _n(a) - _n(b)
        return a

    def mult(self):
        a = self.power()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            b = self.power()
            if op == "*":
                a = _n(a) * _n(b)
            else:
                if _n(b) == 0:
                    raise _XLError("#DIV/0!")
                a = _n(a) / _n(b)
        return a

    def power(self):
        a = self.unary()
        while self.peek()[1] == "^":
            self.take()
            a = _n(a) ** _n(self.unary())
        return a

    def unary(self):
        if self.peek()[1] in ("-", "+"):
            op = self.take()[1]
            v = _n(self.unary())
            return -v if op == "-" else v
        v = self.primary()
        while self.peek()[1] == "%":
            self.take()
            v = _n(v) / 100
        return v

    def primary(self):
        kind, val = self.take()
        if kind == "num":
            return float(val)
        if kind == "str":
            return val[1:-1].replace('""', '"')
        if kind == "bool":
            return val == "TRUE"
        if kind == "err":
            raise _XLError(val)
        if kind == "ref":
            sheet, ref = self.sheet, val
            if "!" in val:
                sheet, ref = val.rsplit("!", 1)
                sheet = sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet
            if ":" in ref:
                a, b = ref.split(":")
                return _Range(self.range(sheet, a, b))
            return self.cell(sheet, ref)
        if kind == "func":
            name = val.upper()
            args = []
            if self.peek()[1] != ")":
                while True:
                    args.append(self.lazy_arg() if name in ("IF", "IFERROR") else self.compare())
                    if self.peek()[1] in (",", ";"):
                        self.take()
                        continue
                    break
            if self.take()[1] != ")":
                raise _Unsupported("missing )")
            return self.call(name, args)
        if val == "(":
            v = self.compare()
            if self.take()[1] != ")":
                raise _Unsupported("missing )")
            return v
        raise _Unsupported(f"unexpected {val!r}")

    def lazy_arg(self):
        """IF/IFERROR arguments: evaluate now, but keep an error as a value."""
        try:
            return self.compare()
        except _XLError as e:
            depth = 0
            while self.i < len(self.toks):            # skip the rest of this argument
                kind, t = self.peek()
                if depth == 0 and t in (",", ";", ")"):
                    break
                if kind == "func" or t == "(":
                    depth += 1
                elif t == ")":
                    depth -= 1
                self.take()
            return e

    def call(self, name: str, args: list):
        nums = lambda: [x for x in _flat(args) if isinstance(x, (int, float)) and not isinstance(x, bool)]  # noqa: E731
        if name == "SUM":
            return float(sum(nums()))
        if name == "AVERAGE":
            xs = nums()
            if not xs:
                raise _XLError("#DIV/0!")
            return sum(xs) / len(xs)
        if name in ("MIN", "MAX"):
            xs = nums()
            return float((min if name == "MIN" else max)(xs)) if xs else 0.0
        if name == "COUNT":
            return float(len(nums()))
        if name == "COUNTA":
            return float(sum(1 for x in _flat(args) if x not in (None, "")))
        if name == "PRODUCT":
            return float(math.prod(nums()))
        if name == "ABS":
            return abs(_n(args[0]))
        if name in ("ROUND", "ROUNDUP", "ROUNDDOWN"):
            from decimal import ROUND_DOWN, ROUND_HALF_UP, ROUND_UP, Decimal

            mode = {"ROUND": ROUND_HALF_UP, "ROUNDUP": ROUND_UP, "ROUNDDOWN": ROUND_DOWN}[name]
            digits = int(_n(args[1])) if len(args) > 1 else 0
            q = Decimal(1).scaleb(-digits)
            return float(Decimal(str(_n(args[0]))).quantize(q, rounding=mode))
        if name == "IF":
            cond = args[0]
            if isinstance(cond, _XLError):
                raise cond
            pick = args[1] if _truthy(cond) else (args[2] if len(args) > 2 else False)
            if isinstance(pick, _XLError):
                raise pick
            return pick
        if name == "IFERROR":
            if isinstance(args[0], _XLError):
                return args[1]
            return args[0]
        if name in ("SUMIF", "COUNTIF", "AVERAGEIF"):
            rng = _flat([args[0]])
            match = _criteria(args[1])
            target = _flat([args[2]]) if len(args) > 2 else rng
            hits = [t for v, t in zip(rng, target) if match(v)]
            if name == "COUNTIF":
                return float(len(hits))
            hn = [h for h in hits if isinstance(h, (int, float)) and not isinstance(h, bool)]
            if name == "SUMIF":
                return float(sum(hn))
            if not hn:
                raise _XLError("#DIV/0!")
            return sum(hn) / len(hn)
        raise _Unsupported(f"function {name}")


class _Range(list):
    pass


def _flat(args) -> list:
    out = []
    for a in args:
        out.extend(a if isinstance(a, list) else [a])
    return out


def _n(v) -> float:
    if isinstance(v, _Range):
        raise _Unsupported("range used as a single value")
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        raise _XLError("#VALUE!")


def _s(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def _truthy(v) -> bool:
    if isinstance(v, str):
        return v.upper() == "TRUE"
    return bool(_n(v))


def _criteria(c):
    import fnmatch

    if isinstance(c, (int, float)) and not isinstance(c, bool):
        return lambda v: isinstance(v, (int, float)) and float(v) == float(c)
    text = str(c)
    m = re.match(r"^(<>|<=|>=|<|>|=)?(.*)$", text)
    op, rhs = m.group(1) or "=", m.group(2)
    try:
        num = float(rhs)
        cmp = {"=": lambda v: v == num, "<>": lambda v: v != num, "<": lambda v: v < num,
               ">": lambda v: v > num, "<=": lambda v: v <= num, ">=": lambda v: v >= num}[op]
        return lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and cmp(float(v))
    except ValueError:
        pat = rhs.lower()
        if op == "<>":
            return lambda v: not fnmatch.fnmatchcase(str(v or "").lower(), pat)
        return lambda v: fnmatch.fnmatchcase(str(v or "").lower(), pat)


def formula_report(path: str, max_values: int = 2000) -> dict:
    """Evaluate every formula in a workbook: how many we could verify, which
    give an Excel error, which we cannot evaluate (unsupported function) and
    their values — to compare with the numbers computed in Python."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    ev = _Evaluator(wb)
    report = {"formulas": 0, "verified": 0, "errors": [], "unverified": [], "values": {}}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                report["formulas"] += 1
                name = f"{ws.title}!{c.coordinate}"
                try:
                    v = ev.cell(ws.title, c.coordinate)
                    report["verified"] += 1
                    if len(report["values"]) < max_values:
                        report["values"][name] = v
                except _XLError as e:
                    report["errors"].append({"cell": name, "error": e.code})
                except (_Unsupported, RecursionError, KeyError, ValueError, TypeError) as e:
                    report["unverified"].append({"cell": name, "reason": str(e)[:80]})
    return report


# ═════════════════════════════════════════════════════════════ file checks ══

def _zip_counts(path: str, prefixes: dict[str, str]) -> dict[str, int]:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return {k: sum(1 for n in names if n.startswith(p) and not n.endswith("/")) for k, p in prefixes.items()}


def _docx_fonts_ok(doc) -> bool:
    xml = doc.styles.element.xml + "".join(p._p.xml for p in doc.paragraphs[:200])
    names = {m.lower() for m in re.findall(r'w:cs="([^"]+)"', xml)}
    return any(n in THAI_FONTS for n in names)


def inspect_docx(path: str) -> dict:
    from docx import Document

    doc = Document(path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    cell_text = sum(1 for t in doc.tables for r in t.rows for c in r.cells if c.text.strip())
    footer_xml = "".join(s.footer._element.xml for s in doc.sections)
    stats = {
        "paragraphs": len(texts),
        "headings": sum(1 for p in doc.paragraphs if p.style is not None and p.text.strip()
                        and p.style.name.lower().startswith(("heading", "title"))),
        "tables": len(doc.tables),
        "images": _zip_counts(path, {"m": "word/media/"})["m"],
        "sections": len(doc.sections),
        "page_numbers": "PAGE" in footer_xml,
        "thai": any(_THAI.search(t) for t in texts),
    }
    issues, warnings = [], []
    if not texts and not cell_text:
        issues.append("เอกสารไม่มีเนื้อหา")
    if stats["thai"] and not _docx_fonts_ok(doc):
        warnings.append("ยังไม่ได้ตั้งฟอนต์ภาษาไทย (complex script) — ตัวอักษรไทยอาจแสดงผลไม่สวย")
    return {"type": "docx", "stats": stats, "issues": issues, "warnings": warnings}


def _shape_kind(shape) -> str:
    try:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return "text"
        if getattr(shape, "has_table", False):
            return "table"
        if getattr(shape, "has_chart", False):
            return "chart"
        if "PICTURE" in str(shape.shape_type):
            return "picture"
    except (NotImplementedError, AttributeError, ValueError):
        pass
    return "other"


def _text_height_pt(shape, default_size: float = 18) -> float:
    """Estimated height of a text frame's text (points)."""
    width = shape.width / EMU_PER_PT - 14              # default insets 0.1"
    if width <= 0:
        return 0.0
    total = 0.0
    for p in shape.text_frame.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size] or ([p.font.size.pt] if p.font.size else [])
        size = max(sizes) if sizes else default_size
        total += estimate_lines([p.text or " "], size, width) * size * 1.2
    return total


def _overflow(shape, default_size: float = 18) -> bool:
    tf = shape.text_frame
    if tf.auto_size is not None and "TEXT_TO_FIT" in str(tf.auto_size):
        return False                                   # PowerPoint shrinks it
    height = shape.height / EMU_PER_PT - 7
    if shape.width / EMU_PER_PT - 14 <= 0 or height <= 0:
        return False
    return _text_height_pt(shape, default_size) > height * 1.15


def _footprint(shape, kind: str) -> tuple[int, int, int, int] | None:
    """The area a shape visibly covers (EMU): a plain text box only covers its
    text, at its vertical anchor; anything else its whole frame."""
    if None in (shape.left, shape.top, shape.width, shape.height) or not shape.width or not shape.height:
        return None
    left, top, right, bottom = shape.left, shape.top, shape.left + shape.width, shape.top + shape.height
    if kind == "text" and "TEXT_BOX" in str(shape.shape_type):
        used = min(shape.height, int((_text_height_pt(shape) + 7) * EMU_PER_PT))
        anchor = str(shape.text_frame.vertical_anchor or "")
        if "MIDDLE" in anchor:
            top += (shape.height - used) // 2
        elif "BOTTOM" in anchor:
            top = bottom - used
        bottom = top + used
    return left, top, right, bottom


def _overlapping(boxes: list[tuple[int, int, int, int]]) -> bool:
    """Any two content areas overlapping by more than 15% of the smaller one."""
    for i, a in enumerate(boxes):
        for b in boxes[i + 1:]:
            w, h = min(a[2], b[2]) - max(a[0], b[0]), min(a[3], b[3]) - max(a[1], b[1])
            smaller = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1]))
            if w > 0 and h > 0 and smaller and w * h > 0.15 * smaller:
                return True
    return False


def inspect_pptx(path: str) -> dict:
    from pptx import Presentation

    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    stats = {"slides": len(prs.slides), "text_boxes": 0, "pictures": 0, "tables": 0, "charts": 0}
    issues, warnings = [], []
    for n, slide in enumerate(prs.slides, 1):
        kinds, boxes = [], []
        for shape in slide.shapes:
            k = _shape_kind(shape)
            kinds.append(k)
            box = _footprint(shape, k) if k in ("text", "table", "chart", "picture") else None
            if box and not (k == "picture" and shape.width * shape.height > 0.6 * sw * sh):   # a background
                boxes.append(box)
            if k == "text":
                stats["text_boxes"] += 1
                if _overflow(shape):
                    warnings.append(f"สไลด์ {n}: ข้อความอาจล้นกรอบ")
            stats["pictures"] += k == "picture"
            stats["tables"] += k == "table"
            stats["charts"] += k == "chart"
            if shape.width and shape.height and shape.left is not None and shape.top is not None and (
                shape.left + shape.width > sw * 1.02 or shape.top + shape.height > sh * 1.02
                or shape.left < -sw * 0.02 or shape.top < -sh * 0.02
            ):
                warnings.append(f"สไลด์ {n}: มีองค์ประกอบล้นขอบสไลด์")
        if _overlapping(boxes):
            warnings.append(f"สไลด์ {n}: ข้อความ/รูป/ตาราง/กราฟวางซ้อนทับกัน")
        pics =[s for s in slide.shapes if _shape_kind(s) == "picture"]
        if pics and not any(k in ("text", "table", "chart") for k in kinds) and \
                any(p.width * p.height > 0.8 * sw * sh for p in pics):
            issues.append(f"สไลด์ {n}: เป็นภาพทั้งหน้า แก้ไขเนื้อหาไม่ได้")
        if not any(k in ("text", "table", "chart", "picture") for k in kinds):
            warnings.append(f"สไลด์ {n}: ว่าง")
    if not stats["slides"]:
        issues.append("ไม่มีสไลด์")
    return {"type": "pptx", "stats": stats, "issues": issues, "warnings": sorted(set(warnings))}


def inspect_xlsx(path: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    parts = _zip_counts(path, {"charts": "xl/charts/chart", "images": "xl/media/"})
    stats = {
        "sheets": wb.sheetnames,
        "cells": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None),
        "charts": parts["charts"],
        "images": parts["images"],
    }
    issues, warnings = [], []
    rep = formula_report(path)
    stats.update(formulas=rep["formulas"], formulas_verified=rep["verified"])
    if rep["errors"]:
        issues.extend(f"สูตรผิดพลาด {e['cell']} → {e['error']}" for e in rep["errors"][:10])
    if rep["unverified"]:
        warnings.append(f"ตรวจผลสูตรไม่ได้ {len(rep['unverified'])} ช่อง (ฟังก์ชันที่ระบบตรวจไม่รองรับ)")
    if not stats["cells"]:
        issues.append("ไฟล์ Excel ว่าง")
    return {"type": "xlsx", "stats": stats, "issues": issues, "warnings": warnings, "formulas": rep}


def inspect_pdf(path: str) -> dict:
    stats: dict = {}
    issues, warnings = [], []
    try:
        import fitz

        with fitz.open(path) as doc:
            if doc.needs_pass:
                return {"type": "pdf", "stats": {}, "issues": ["PDF มีรหัสผ่าน"], "warnings": []}
            texts = [p.get_text() for p in doc]
            fonts = {f[3] for p in doc for f in p.get_fonts()}
            stats = {"pages": doc.page_count, "text_chars": sum(len(t.strip()) for t in texts),
                     "images": sum(len(p.get_images()) for p in doc), "fonts": sorted(fonts)[:20],
                     "thai": any(_THAI.search(t) for t in texts)}
    except ImportError:
        from pypdf import PdfReader

        r = PdfReader(path)
        texts = [(pg.extract_text() or "") for pg in r.pages]
        stats = {"pages": len(r.pages), "text_chars": sum(len(t.strip()) for t in texts),
                 "thai": any(_THAI.search(t) for t in texts)}
    if not stats.get("pages"):
        issues.append("PDF ไม่มีหน้า")
    elif not stats.get("text_chars"):
        warnings.append("PDF ไม่มีข้อความที่คัดลอก/ค้นหาได้ (อาจเป็นภาพสแกน)")
    return {"type": "pdf", "stats": stats, "issues": issues, "warnings": warnings}


_INSPECTORS = {".docx": inspect_docx, ".pptx": inspect_pptx, ".xlsx": inspect_xlsx, ".pdf": inspect_pdf}


def _comparable(path: str) -> dict:
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        from docx import Document

        d = Document(path)
        return {"ตาราง": len(d.tables), "รูปภาพ": _zip_counts(path, {"m": "word/media/"})["m"],
                "ย่อหน้าที่มีข้อความ": sum(1 for p in d.paragraphs if p.text.strip())}
    if ext == ".pptx":
        info = inspect_pptx(path)["stats"]
        return {"สไลด์": info["slides"], "รูปภาพ": info["pictures"], "ตาราง": info["tables"],
                "กราฟ": info["charts"]}
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=False)
        parts = _zip_counts(path, {"c": "xl/charts/chart", "i": "xl/media/", "p": "xl/pivotTables/",
                                   "m": "xl/comments", "d": "xl/drawings/"})
        out = {"ชีต": len(wb.sheetnames), "กราฟ": parts["c"], "รูปภาพ": parts["i"], "Pivot Table": parts["p"],
               "ความคิดเห็น": parts["m"], "Defined names": len(list(wb.defined_names))}
        out["สูตร"] = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                          if isinstance(c.value, str) and c.value.startswith("="))
        out["_sheets"] = wb.sheetnames
        return out
    return {}


def compare(original: str, edited: str) -> list[str]:
    """What the edited file lost compared with the original (same type)."""
    a, b = _comparable(original), _comparable(edited)
    losses = []
    for key, before in a.items():
        if key.startswith("_"):
            continue
        after = b.get(key, 0)
        if key == "ย่อหน้าที่มีข้อความ":
            if before >= 5 and after < before * 0.7:
                losses.append(f"ข้อความหายไปมาก ({before} → {after} ย่อหน้า)")
        elif after < before:
            losses.append(f"{key} ลดลงจาก {before} เหลือ {after}")
    missing = [s for s in a.get("_sheets", []) if s not in b.get("_sheets", [])]
    if missing:
        losses.append("ชีตหายไป: " + ", ".join(missing))
    return losses


def check(path: str, original: str | None = None) -> dict:
    """Re-open the file and check it. {"ok", "type", "stats", "issues",
    "warnings"} — issues mean the file should not be delivered as is."""
    ext = Path(path).suffix.lower()
    inspector = _INSPECTORS.get(ext)
    if inspector is None:
        return {"ok": True, "type": ext.lstrip("."), "stats": {}, "issues": [], "warnings": []}
    try:
        if ext in (".docx", ".pptx", ".xlsx"):
            with zipfile.ZipFile(path) as z:
                bad = z.testzip()
            if bad:
                raise ValueError(f"damaged part {bad}")
        result = inspector(path)
    except Exception as e:  # noqa: BLE001 — a file we cannot open is the finding
        return {"ok": False, "type": ext.lstrip("."), "stats": {},
                "issues": [f"เปิดไฟล์ไม่ได้ ({type(e).__name__}: {str(e)[:120]})"], "warnings": []}
    if original and Path(original).suffix.lower() == ext:
        try:
            result["warnings"] = result["warnings"] + compare(original, path)
        except Exception:  # noqa: BLE001
            pass
    result.pop("formulas", None)
    result["ok"] = not result["issues"]
    return result

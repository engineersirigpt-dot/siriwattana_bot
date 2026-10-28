"""
AI usage log — tokens and cost of every billed OpenAI use, for the admin usage
report and its comparison with paying for ChatGPT Business seats instead.

Rows are written where the cost is already computed (one source of truth):
  chat_pg.save_chat_message_pg     → the answer (feature from its source)
  chat_pg.add_usage_to_message_pg  → the safety classifiers of an answered turn
  main (blocked turns)             → the safety classifiers of a blocked turn
  doc_translation._run             → a finished translation job
The table (db_pg: ai_usage_log) is append-only and independent of
chat_history, so deleting a chat keeps its cost. It holds no question text.

Not logged (no usage returned / negligible): audio transcription, the RAG
embedding of each question, the image-intent and image-prompt helper calls.
"""

from __future__ import annotations

import io
import os
from datetime import date, datetime, timedelta, timezone

# ChatGPT Business price per user per month (USD) used as the comparison
# default; the admin page lets the reader change it. Check the current price.
CHATGPT_BUSINESS_SEAT_USD = float(os.getenv("CHATGPT_BUSINESS_SEAT_USD", "30"))
USD_TO_THB = float(os.getenv("DASHBOARD_USD_TO_THB", "36"))
RAW_EXPORT_LIMIT = 50_000

_BANGKOK = timezone(timedelta(hours=7))

# chat_history.source → report feature
_FEATURE_OF_SOURCE = {
    "llm": "chat", "llm-calc": "chat", "rag": "chat", "rag-calc": "chat",
    "brain": "chat", "brain-calc": "chat",
    "files": "files",
    "image_gen": "image", "image_edit": "image",
    "web": "web_search", "url": "read_link", "agent": "agent",
    "code": "code_interpreter", "code_limit": "code_interpreter", "document": "document_engine",
    "linsip": "linsip",
}
FEATURE_LABELS = {
    "chat": "แชททั่วไป / ฐานความรู้",
    "files": "อ่านไฟล์ / ตรวจคำผิด",
    "image": "สร้าง / แก้รูป",
    "web_search": "ค้นเว็บ",
    "read_link": "อ่านลิงก์",
    "agent": "Agent (ค้นเว็บหลายขั้น)",
    "code_interpreter": "Code Interpreter (Python)",
    "document_engine": "Document Engine (สร้าง/แก้เอกสาร)",
    "linsip": "ค้นงาน linsip",
    "translation": "แปลเอกสาร",
    "safety": "ตัวตรวจข้อมูลอ่อนไหว",
}


def feature_for(source: str | None) -> str:
    return _FEATURE_OF_SOURCE.get(source or "", source or "other")


def _usage_fields(usage: dict) -> tuple[str | None, int, int, float]:
    """Accepts both shapes: llm.accumulate_usage (model_used) and per-call (model)."""
    return (
        usage.get("model_used") or usage.get("model"),
        int(usage.get("prompt_tokens") or 0),
        int(usage.get("completion_tokens") or 0),
        float(usage.get("cost_usd") or 0.0),
    )


def insert_usage(
    cur, *, user_id: int | None, feature: str, source: str | None, usage: dict,
    session_id: int | None = None, message_id: int | None = None, ref: str | None = None,
    created_at: datetime | None = None, backfilled: bool = False,
) -> None:
    """Write one log row with the caller's cursor (same transaction)."""
    model, p, c, cost = _usage_fields(usage)
    cur.execute(
        """
        INSERT INTO ai_usage_log
            (created_at, user_id, username, feature, source, model,
             prompt_tokens, completion_tokens, cost_usd,
             session_id, message_id, ref, backfilled)
        VALUES (COALESCE(%s, now()), %s, (SELECT username FROM users WHERE id = %s), %s, %s, %s,
                %s, %s, %s, %s, %s, %s, %s)
        """,
        (created_at, user_id, user_id, feature, source, model, p, c, cost,
         session_id, message_id, ref, backfilled),
    )


def insert_usage_safe(conn, cur, **kw) -> None:
    """insert_usage inside a SAVEPOINT of the caller's transaction: a logging
    problem rolls back only the log row, never the answer being saved."""
    try:
        with conn.transaction():
            insert_usage(cur, **kw)
    except Exception as e:  # noqa: BLE001
        import sys

        print(f"USAGE_LOG_FAILED: {e!r}", file=sys.stderr)


def log_usage_pg(
    user_id: int | None, feature: str, source: str | None, usage: dict, *, ref: str | None = None,
) -> None:
    """Write one log row on its own connection. Never raises: a logging
    problem must not break the reply the user is waiting for."""
    if not usage:
        return
    try:
        from db_pg import get_pg_conn

        with get_pg_conn() as conn, conn.cursor() as cur:
            insert_usage(cur, user_id=user_id, feature=feature, source=source, usage=usage, ref=ref)
    except Exception as e:  # noqa: BLE001
        import sys

        print(f"USAGE_LOG_FAILED: {e!r}", file=sys.stderr)


# ── Report ──────────────────────────────────────────────────────────────────

def _is_request(feature: str, source: str | None) -> bool:
    """A user request. The safety classifiers of an ANSWERED turn are an add-on
    to that request; a BLOCKED turn is a request of its own."""
    return feature != "safety" or source == "blocked"


def months_between(since: date, until: date) -> list[str]:
    out, y, m = [], since.year, since.month
    while (y, m) <= (until.year, until.month):
        out.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def build_report(rows: list[dict], *, since: date, until: date, seat_usd: float,
                 usd_to_thb: float = USD_TO_THB) -> dict:
    """Aggregate grouped log rows into the report. Pure — unit-tested.

    rows: [{month, user_id, username, feature, source, model, n,
            prompt_tokens, completion_tokens, cost_usd}] (n = log rows).
    ChatGPT Business comparison: every user active in a month would need a
    seat that month (seat_usd × active users); per user, SiriChat's average
    cost per active month vs one seat."""
    def blank() -> dict:
        return {"requests": 0, "prompt_tokens": 0, "completion_tokens": 0, "cost_usd": 0.0}

    def add(acc: dict, r: dict) -> None:
        if _is_request(r["feature"], r.get("source")):
            acc["requests"] += int(r["n"])
        acc["prompt_tokens"] += int(r["prompt_tokens"])
        acc["completion_tokens"] += int(r["completion_tokens"])
        acc["cost_usd"] += float(r["cost_usd"])

    def finish(acc: dict) -> dict:
        acc["total_tokens"] = acc["prompt_tokens"] + acc["completion_tokens"]
        acc["cost_usd"] = round(acc["cost_usd"], 6)
        acc["cost_thb"] = round(acc["cost_usd"] * usd_to_thb, 2)
        return acc

    months = months_between(since, until)
    monthly = {m: {**blank(), "users": set()} for m in months}
    users: dict = {}
    features: dict = {}
    models: dict = {}
    total = blank()
    for r in rows:
        m = monthly.setdefault(r["month"], {**blank(), "users": set()})
        add(m, r)
        add(total, r)
        if _is_request(r["feature"], r.get("source")) and r.get("user_id") is not None:
            m["users"].add(r["user_id"])
        key = r.get("user_id")
        u = users.setdefault(key, {**blank(), "username": r.get("username") or str(key), "months": set()})
        add(u, r)
        if _is_request(r["feature"], r.get("source")):
            u["months"].add(r["month"])
        f = features.setdefault(r["feature"], {**blank(), "feature": r["feature"],
                                               "label": FEATURE_LABELS.get(r["feature"], r["feature"])})
        add(f, r)
        md = models.setdefault(r.get("model") or "(ไม่ระบุ)", {**blank(), "model": r.get("model") or "(ไม่ระบุ)"})
        add(md, r)

    seat_thb = seat_usd * usd_to_thb
    monthly_out = []
    for month in sorted(monthly):
        m = monthly[month]
        active = len(m.pop("users"))
        row = finish({"month": month, "active_users": active, **m})
        row["chatgpt_usd"] = round(active * seat_usd, 2)
        row["chatgpt_thb"] = round(active * seat_thb, 2)
        row["difference_thb"] = round(row["chatgpt_thb"] - row["cost_thb"], 2)   # > 0: SiriChat cheaper
        monthly_out.append(row)

    users_out = []
    for u in users.values():
        months_active = len(u.pop("months"))
        row = finish(u)
        row["months_active"] = months_active
        row["avg_month_thb"] = round(row["cost_thb"] / months_active, 2) if months_active else 0.0
        row["seat_thb"] = round(seat_thb, 2)
        row["siri_cheaper"] = row["avg_month_thb"] <= seat_thb
        users_out.append(row)
    users_out.sort(key=lambda u: -u["cost_usd"])

    totals = finish(total)
    totals["active_users"] = len({k for k, u in users.items() if u.get("months_active")})
    totals["chatgpt_usd"] = round(sum(m["chatgpt_usd"] for m in monthly_out), 2)
    totals["chatgpt_thb"] = round(sum(m["chatgpt_thb"] for m in monthly_out), 2)
    totals["difference_thb"] = round(totals["chatgpt_thb"] - totals["cost_thb"], 2)
    per_request = totals["cost_usd"] / totals["requests"] if totals["requests"] else 0.0
    totals["avg_cost_per_request_thb"] = round(per_request * usd_to_thb, 4)
    # requests a user must make per month before one seat would be cheaper
    totals["break_even_requests_per_month"] = int(seat_usd / per_request) if per_request else None

    return {
        "period": {"since": since.isoformat(), "until": until.isoformat(), "months": months},
        "seat_usd": seat_usd,
        "seat_thb": round(seat_thb, 2),
        "usd_to_thb": usd_to_thb,
        "totals": totals,
        "monthly": monthly_out,
        "users": users_out,
        "features": sorted((finish(f) for f in features.values()), key=lambda f: -f["cost_usd"]),
        "models": sorted((finish(m) for m in models.values()), key=lambda m: -m["cost_usd"]),
    }


_RANGE_SQL = (
    "created_at >= (%s::date)::timestamp AT TIME ZONE 'Asia/Bangkok' "
    "AND created_at < ((%s::date) + 1)::timestamp AT TIME ZONE 'Asia/Bangkok'"
)


def usage_report_pg(since: date, until: date, seat_usd: float) -> dict:
    """Report for [since, until] (Bangkok calendar days, both inclusive)."""
    from db_pg import get_pg_conn

    with get_pg_conn() as conn, conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT to_char(created_at AT TIME ZONE 'Asia/Bangkok', 'YYYY-MM') AS month,
                   user_id, max(username), feature, source, model, count(*),
                   sum(prompt_tokens), sum(completion_tokens), sum(cost_usd)
            FROM ai_usage_log
            WHERE {_RANGE_SQL}
            GROUP BY 1, 2, 4, 5, 6
            """,
            (since, until),
        )
        keys = ("month", "user_id", "username", "feature", "source", "model", "n",
                "prompt_tokens", "completion_tokens", "cost_usd")
        rows = [dict(zip(keys, r)) for r in cur.fetchall()]
    return build_report(rows, since=since, until=until, seat_usd=seat_usd)


def usage_rows_pg(since: date, until: date, limit: int = RAW_EXPORT_LIMIT) -> list[tuple]:
    from db_pg import get_pg_conn

    with get_pg_conn() as conn, conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT created_at AT TIME ZONE 'Asia/Bangkok', username, feature, source, model,
                   prompt_tokens, completion_tokens, cost_usd, backfilled
            FROM ai_usage_log
            WHERE {_RANGE_SQL}
            ORDER BY created_at
            LIMIT %s
            """,
            (since, until, limit),
        )
        return cur.fetchall()


def build_usage_xlsx(report: dict, raw_rows: list[tuple]) -> bytes:
    """The report as an Excel workbook for sending on (Thai headers)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fx = report["usd_to_thb"]
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="7C3AED")
    wb = Workbook()

    def sheet(title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "A2"
        return ws

    t = report["totals"]
    wb.remove(wb.active)
    sheet("สรุป", ["หัวข้อ", "ค่า"], [
        ["ช่วงเวลา", f"{report['period']['since']} ถึง {report['period']['until']}"],
        ["อัตราแลกเปลี่ยน (บาท/USD)", fx],
        ["ราคา ChatGPT Business ต่อคนต่อเดือน (USD)", report["seat_usd"]],
        ["จำนวนคำขอ (request)", t["requests"]],
        ["ผู้ใช้ที่ใช้งาน", t["active_users"]],
        ["Token ขาเข้า (input)", t["prompt_tokens"]],
        ["Token ขาออก (output)", t["completion_tokens"]],
        ["Token รวม", t["total_tokens"]],
        ["ค่าใช้จ่าย SiriChat (USD)", t["cost_usd"]],
        ["ค่าใช้จ่าย SiriChat (บาท)", t["cost_thb"]],
        ["ถ้าใช้ ChatGPT Business แทน (บาท)", t["chatgpt_thb"]],
        ["ส่วนต่าง (บาท, บวก = SiriChat ถูกกว่า)", t["difference_thb"]],
        ["ค่าใช้จ่ายเฉลี่ยต่อคำขอ (บาท)", t["avg_cost_per_request_thb"]],
        ["คำขอ/คน/เดือน ที่ ChatGPT Business จะเริ่มคุ้มกว่า", t["break_even_requests_per_month"]],
        [],
        ["หมายเหตุ", "ChatGPT Business คิดเป็นรายที่นั่งต่อเดือนไม่ว่าจะใช้มากหรือน้อย; "
                     "SiriChat คิดตามการใช้งานจริง (token) ผ่าน OpenAI API"],
        ["", "ไม่รวม: ค่าถอดเสียง, embedding ของการค้นฐานความรู้ และการเรียก AI ช่วยเล็ก ๆ "
             "ของการสร้างรูป (มีค่าน้อยมาก)"],
        ["", "แถว backfill = ข้อมูลย้อนหลังที่ดึงจากระบบก่อนมี log นี้ (จำนวน token อาจไม่ครบ)"],
    ], [52, 70])

    sheet("รายเดือน", ["เดือน", "ผู้ใช้ที่ใช้งาน", "คำขอ", "Token ขาเข้า", "Token ขาออก", "Token รวม",
                       "ค่าใช้จ่าย (USD)", "ค่าใช้จ่าย (บาท)", "ChatGPT Business (บาท)", "ส่วนต่าง (บาท)"],
          [[m["month"], m["active_users"], m["requests"], m["prompt_tokens"], m["completion_tokens"],
            m["total_tokens"], m["cost_usd"], m["cost_thb"], m["chatgpt_thb"], m["difference_thb"]]
           for m in report["monthly"]])
    sheet("รายผู้ใช้", ["ผู้ใช้", "เดือนที่ใช้งาน", "คำขอ", "Token ขาเข้า", "Token ขาออก", "Token รวม",
                        "ค่าใช้จ่าย (บาท)", "เฉลี่ยต่อเดือน (บาท)", "ChatGPT Business ต่อเดือน (บาท)",
                        "แบบไหนถูกกว่า"],
          [[u["username"], u["months_active"], u["requests"], u["prompt_tokens"], u["completion_tokens"],
            u["total_tokens"], u["cost_thb"], u["avg_month_thb"], u["seat_thb"],
            "SiriChat" if u["siri_cheaper"] else "ChatGPT Business"] for u in report["users"]])
    sheet("ตามฟีเจอร์", ["ฟีเจอร์", "คำขอ", "Token ขาเข้า", "Token ขาออก", "ค่าใช้จ่าย (USD)", "ค่าใช้จ่าย (บาท)"],
          [[f["label"], f["requests"], f["prompt_tokens"], f["completion_tokens"], f["cost_usd"], f["cost_thb"]]
           for f in report["features"]], [32, 12, 16, 16, 16, 16])
    sheet("ตามโมเดล", ["โมเดล", "คำขอ", "Token ขาเข้า", "Token ขาออก", "ค่าใช้จ่าย (USD)", "ค่าใช้จ่าย (บาท)"],
          [[m["model"], m["requests"], m["prompt_tokens"], m["completion_tokens"], m["cost_usd"], m["cost_thb"]]
           for m in report["models"]])
    sheet("รายการทั้งหมด", ["เวลา", "ผู้ใช้", "ฟีเจอร์", "ประเภท", "โมเดล", "Token ขาเข้า", "Token ขาออก",
                            "ค่าใช้จ่าย (USD)", "ค่าใช้จ่าย (บาท)", "ข้อมูลย้อนหลัง"],
          [[ts.replace(tzinfo=None) if hasattr(ts, "replace") else ts, user,
            FEATURE_LABELS.get(feature, feature), source, model, p, c,
            float(cost), round(float(cost) * fx, 4), "ใช่" if back else ""]
           for ts, user, feature, source, model, p, c, cost, back in raw_rows],
          [20, 14, 28, 14, 18, 14, 14, 16, 16, 14])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

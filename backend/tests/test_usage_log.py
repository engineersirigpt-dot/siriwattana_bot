"""
Tests for the AI usage log (usage_log.py): what gets logged where, the report
maths behind the ChatGPT Business comparison, the Excel export and the admin
endpoints. Offline — postgres behaviour is in tests/test_usage_log_pg.py.
"""

import io
from datetime import date
from types import SimpleNamespace as NS

import openpyxl
import pytest
from starlette.testclient import TestClient

import doc_translation
import main
import usage_log as ul

ADMIN = {"id": 1, "role": "admin", "username": "2690028"}
USER = {"id": 2, "role": "user", "username": "2133557"}


# ── feature mapping ─────────────────────────────────────────────────────────

@pytest.mark.parametrize("source,feature", [
    ("llm", "chat"), ("rag-calc", "chat"), ("brain", "chat"), ("files", "files"),
    ("image_gen", "image"), ("image_edit", "image"), ("web", "web_search"),
    ("agent", "agent"), ("code", "code_interpreter"), ("code_limit", "code_interpreter"),
    ("linsip", "linsip"), ("url", "read_link"), ("something_new", "something_new"), (None, "other"),
])
def test_feature_for(source, feature):
    assert ul.feature_for(source) == feature


# ── report maths ────────────────────────────────────────────────────────────

def _row(month, uid, feature, source, n, p, c, usd, model="gpt-4.1", name=None):
    return dict(month=month, user_id=uid, username=name or f"u{uid}", feature=feature, source=source,
                model=model, n=n, prompt_tokens=p, completion_tokens=c, cost_usd=usd)


ROWS = [
    # user 1: heavy in August (Code Interpreter), light in September
    _row("2026-08", 1, "chat", "llm", 40, 80_000, 8_000, 0.40),
    _row("2026-08", 1, "safety", "classifier", 40, 40_000, 1_000, 0.006, "gpt-4o-mini"),  # add-on, not requests
    _row("2026-08", 1, "code_interpreter", "code", 100, 600_000, 300_000, 9.0, "gpt-5"),
    _row("2026-09", 1, "chat", "llm", 10, 20_000, 2_000, 0.10),
    # user 2: light, one blocked question (a request of its own)
    _row("2026-08", 2, "chat", "llm", 5, 10_000, 1_000, 0.05),
    _row("2026-08", 2, "safety", "blocked", 1, 900, 20, 0.0002, "gpt-4o-mini"),
    _row("2026-09", 2, "translation", "translation", 1, 50_000, 40_000, 0.42),
]


def test_report_totals_and_requests():
    r = ul.build_report(ROWS, since=date(2026, 7, 1), until=date(2026, 9, 30), seat_usd=30, usd_to_thb=36)
    t = r["totals"]
    # classifier rows of answered turns are not extra requests; a blocked turn is
    assert t["requests"] == 40 + 100 + 10 + 5 + 1 + 1
    assert t["prompt_tokens"] == sum(x["prompt_tokens"] for x in ROWS)
    assert t["total_tokens"] == t["prompt_tokens"] + t["completion_tokens"]
    assert t["cost_usd"] == pytest.approx(9.9762)
    assert t["cost_thb"] == pytest.approx(9.9762 * 36, abs=0.01)
    assert t["active_users"] == 2
    assert r["period"]["months"] == ["2026-07", "2026-08", "2026-09"]


def test_monthly_comparison_with_chatgpt_business_seats():
    r = ul.build_report(ROWS, since=date(2026, 7, 1), until=date(2026, 9, 30), seat_usd=30, usd_to_thb=36)
    jul, aug, sep = r["monthly"]
    assert jul["active_users"] == 0 and jul["chatgpt_thb"] == 0       # empty month still listed
    assert aug["active_users"] == 2 and aug["chatgpt_thb"] == 2 * 30 * 36
    assert aug["cost_thb"] == pytest.approx(9.4562 * 36, abs=0.01)
    assert aug["difference_thb"] == pytest.approx(2160 - 9.4562 * 36, abs=0.01)   # SiriChat cheaper
    assert sep["active_users"] == 2
    assert r["totals"]["chatgpt_thb"] == pytest.approx(4 * 30 * 36)   # 2 users × 2 active months


def test_per_user_average_vs_one_seat():
    r = ul.build_report(ROWS, since=date(2026, 7, 1), until=date(2026, 9, 30), seat_usd=3, usd_to_thb=36)
    u1, u2 = r["users"]                                               # sorted by cost
    assert u1["username"] == "u1" and u1["months_active"] == 2
    assert u1["avg_month_thb"] == pytest.approx((9.506 * 36) / 2, abs=0.01)
    assert u1["siri_cheaper"] is False                                # heavy user: a $3 seat would win
    assert u2["siri_cheaper"] is True


def test_break_even_and_average_per_request():
    r = ul.build_report(ROWS, since=date(2026, 7, 1), until=date(2026, 9, 30), seat_usd=30, usd_to_thb=36)
    per_req = 9.9762 / 157
    assert r["totals"]["avg_cost_per_request_thb"] == pytest.approx(per_req * 36, abs=1e-3)
    assert r["totals"]["break_even_requests_per_month"] == int(30 / per_req)


def test_features_and_models_breakdown():
    r = ul.build_report(ROWS, since=date(2026, 8, 1), until=date(2026, 9, 30), seat_usd=30)
    feats = {f["feature"]: f for f in r["features"]}
    assert feats["code_interpreter"]["label"] == "Code Interpreter (Python)"
    assert feats["safety"]["requests"] == 1                           # only the blocked one
    assert r["features"][0]["feature"] == "code_interpreter"          # most expensive first
    assert {m["model"] for m in r["models"]} == {"gpt-4.1", "gpt-4o-mini", "gpt-5"}


def test_empty_period():
    r = ul.build_report([], since=date(2026, 10, 1), until=date(2026, 10, 31), seat_usd=30)
    assert r["totals"]["requests"] == 0 and r["totals"]["break_even_requests_per_month"] is None
    assert r["users"] == [] and len(r["monthly"]) == 1


def test_months_between_crosses_years():
    assert ul.months_between(date(2026, 11, 15), date(2027, 2, 1)) == ["2026-11", "2026-12", "2027-01", "2027-02"]


# ── Excel ───────────────────────────────────────────────────────────────────

def test_excel_export_has_every_view():
    from datetime import datetime

    report = ul.build_report(ROWS, since=date(2026, 8, 1), until=date(2026, 9, 30), seat_usd=30)
    raw = [(datetime(2026, 8, 3, 9, 0), "u1", "code_interpreter", "code", "gpt-5", 6000, 3000, 0.09, False),
           (datetime(2026, 7, 1, 10, 0), "u2", "image", "image_gen", "gpt-image-2", 0, 0, 0.16, True)]
    wb = openpyxl.load_workbook(io.BytesIO(ul.build_usage_xlsx(report, raw)))
    assert wb.sheetnames == ["สรุป", "รายเดือน", "รายผู้ใช้", "ตามฟีเจอร์", "ตามโมเดล", "รายการทั้งหมด"]
    summary = {row[0]: row[1] for row in wb["สรุป"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert summary["Token รวม"] == report["totals"]["total_tokens"]
    assert summary["ราคา ChatGPT Business ต่อคนต่อเดือน (USD)"] == 30
    rows = list(wb["รายการทั้งหมด"].iter_rows(min_row=2, values_only=True))
    assert rows[0][2] == "Code Interpreter (Python)" and rows[0][5] == 6000
    assert rows[1][-1] == "ใช่"                                       # backfilled marked


# ── what gets logged ────────────────────────────────────────────────────────

def test_insert_usage_safe_never_breaks_the_caller():
    class Conn:
        def transaction(self):
            class Tx:
                def __enter__(self_inner):
                    return self_inner

                def __exit__(self_inner, *exc):
                    return False
            return Tx()

    class Cur:
        def execute(self, *a):
            raise RuntimeError("table missing")

    ul.insert_usage_safe(Conn(), Cur(), user_id=1, feature="chat", source="llm",
                         usage={"model_used": "gpt-4.1", "cost_usd": 0.01})   # no exception


def test_insert_usage_accepts_both_usage_shapes():
    sent = []
    cur = NS(execute=lambda sql, args: sent.append(args))
    ul.insert_usage(cur, user_id=3, feature="chat", source="llm",
                    usage={"model_used": "gpt-4.1", "prompt_tokens": 10, "completion_tokens": 2, "cost_usd": 0.5})
    ul.insert_usage(cur, user_id=3, feature="safety", source="blocked",
                    usage={"model": "gpt-4o-mini", "prompt_tokens": 900, "completion_tokens": 20, "cost_usd": 0.0002})
    assert sent[0][5:9] == ("gpt-4.1", 10, 2, 0.5)
    assert sent[1][5:9] == ("gpt-4o-mini", 900, 20, 0.0002)


def test_blocked_turn_is_logged_as_its_own_request(monkeypatch):
    logged = []
    monkeypatch.setattr(main, "use_postgres_auth", lambda: True)
    monkeypatch.setattr(ul, "log_usage_pg", lambda *a, **kw: logged.append(a))
    main._log_blocked_usage(USER, [
        {"model": "gpt-4o-mini", "prompt_tokens": 1000, "completion_tokens": 30, "cost_usd": 0.0002},
        {"model": "gpt-4o-mini", "prompt_tokens": 500, "completion_tokens": 20, "cost_usd": 0.0001},
    ])
    uid, feature, source, usage = logged[0]
    assert (uid, feature, source) == (2, "safety", "blocked")
    assert usage["prompt_tokens"] == 1500 and usage["cost_usd"] == pytest.approx(0.0003)
    main._log_blocked_usage(USER, [])                                # keyword hit: no classifier ran
    assert len(logged) == 1


def test_stream_blocked_path_writes_the_log(monkeypatch):
    from tests.test_code_interpreter import CLS_USAGE, SCOPE_USAGE, _stream

    logged = []
    monkeypatch.setattr(ul, "log_usage_pg", lambda *a, **kw: logged.append(a))
    events, _ = _stream(monkeypatch, flagged=("PII", "r"), decision="denied")
    assert events[-1]["source"] == "blocked"
    assert logged[0][1:3] == ("safety", "blocked")
    assert logged[0][3]["cost_usd"] == pytest.approx(CLS_USAGE["cost_usd"] + SCOPE_USAGE["cost_usd"])


def test_finished_translation_job_is_logged(monkeypatch, tmp_path):
    logged = []
    monkeypatch.setenv("DATABASE_URL", "postgresql://unused")         # _db_on; DB writes fail quietly
    monkeypatch.setattr(ul, "log_usage_pg", lambda *a, **kw: logged.append((a, kw)))
    monkeypatch.setattr(doc_translation.tm, "translate_pdf", lambda *a, **kw: {
        "docx": "x.docx", "pdf": None, "review": None, "n_flagged": 0, "n_pages": 3,
        "tokens_in": 12_000, "tokens_out": 9_000, "cost_usd": 0.096})
    doc_translation._jobs["job1"] = {"id": "job1", "user_id": 7, "total": 3}
    try:
        doc_translation._run("job1", str(tmp_path / "a.pdf"), str(tmp_path), "a")
    finally:
        doc_translation._jobs.pop("job1", None)
    (uid, feature, source, usage), kw = logged[0]
    assert (uid, feature, source) == (7, "translation", "translation")
    assert usage["prompt_tokens"] == 12_000 and usage["cost_usd"] == 0.096 and kw["ref"] == "job1"


# ── admin endpoints ─────────────────────────────────────────────────────────

def _get(monkeypatch, path, user):
    monkeypatch.setattr(main, "use_postgres_auth", lambda: True)
    monkeypatch.setattr(main, "audit_log", lambda *a, **kw: None)
    main.app.dependency_overrides[main.current_user] = lambda: user
    try:
        return TestClient(main.app).get(path)
    finally:
        main.app.dependency_overrides.pop(main.current_user, None)


def test_report_endpoint_is_admin_only(monkeypatch):
    monkeypatch.setattr("auth.audit_log", lambda *a, **kw: None)
    assert _get(monkeypatch, "/admin/usage/report", USER).status_code == 403
    assert _get(monkeypatch, "/admin/usage/export.xlsx", USER).status_code == 403


def test_report_endpoint_passes_range_and_seat(monkeypatch):
    seen = []
    monkeypatch.setattr(ul, "usage_report_pg", lambda s, u, seat: seen.append((s, u, seat)) or {"ok": 1})
    r = _get(monkeypatch, "/admin/usage/report?since=2026-07-01&until=2026-10-28&seat_usd=25", ADMIN)
    assert r.status_code == 200 and r.json() == {"ok": 1}
    assert seen == [(date(2026, 7, 1), date(2026, 10, 28), 25.0)]

    _get(monkeypatch, "/admin/usage/report", ADMIN)                 # defaults: this month, env seat
    s, u, seat = seen[1]
    assert s.day == 1 and s <= u and seat == ul.CHATGPT_BUSINESS_SEAT_USD


@pytest.mark.parametrize("q", ["since=2026-13-01", "since=2026-10-28&until=2026-07-01",
                               "since=2020-01-01&until=2026-10-01"])
def test_report_endpoint_rejects_bad_ranges(monkeypatch, q):
    monkeypatch.setattr(ul, "usage_report_pg", lambda *a: {})
    assert _get(monkeypatch, f"/admin/usage/report?{q}", ADMIN).status_code == 400


def test_export_endpoint_returns_a_workbook(monkeypatch):
    report = ul.build_report(ROWS, since=date(2026, 8, 1), until=date(2026, 9, 30), seat_usd=30)
    monkeypatch.setattr(ul, "usage_report_pg", lambda *a: report)
    monkeypatch.setattr(ul, "usage_rows_pg", lambda *a: [])
    r = _get(monkeypatch, "/admin/usage/export.xlsx?since=2026-08-01&until=2026-09-30", ADMIN)
    assert r.status_code == 200
    assert "SiriChat_AI_usage_2026-08-01_2026-09-30.xlsx" in r.headers["content-disposition"]
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames[0] == "สรุป"

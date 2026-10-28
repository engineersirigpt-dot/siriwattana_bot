"""
V4 in the chat pipeline (/chat/stream, /chat, the agent turn, /memory):
routing replaces the classifier only when its flag is on, the routed model and
effort reach the answer call, every answer logs its route; memory commands,
memory + summary in the history, per-user /memory endpoints. Flags off →
exactly the legacy calls. Offline.
"""

import json
from types import SimpleNamespace as NS

import pytest
from starlette.testclient import TestClient

import agent
import chat_pg
import context_manager as cm
import main
import memory
import model_router as mr
from tests.test_memory import FakeStore

ADMIN = {"id": 1, "role": "admin", "username": "a"}
OTHER = {"id": 2, "role": "admin", "username": "b"}
CLS = {"model": "gpt-4o-mini", "prompt_tokens": 60, "completion_tokens": 1, "cost_usd": 0.00001}
HIT = {"id": 3, "question": "ลาพักร้อนกี่วัน", "answer": "10 วัน", "similarity": 0.9, "source_file": None}


@pytest.fixture
def pipe(monkeypatch):
    s = NS(saved=[], classified=[], streamed=[], answered=[], audits=[], summaries=[], hit=None, flag=None,
           safety=[], scoped=[],
           history=[{"role": "user", "content": "ก่อนหน้า"}, {"role": "assistant", "content": "ตอบก่อนหน้า"}],
           store=FakeStore())
    monkeypatch.setattr(mr, "MODEL_ROUTING_MODE", "off")
    monkeypatch.setattr(memory, "MEMORY_MODE", "off")
    monkeypatch.setattr(mr, "TIERS", {"light": ("gpt-5-mini", "minimal"), "standard": ("gpt-5", "low"),
                                      "calc": ("gpt-5-mini", "medium"), "deep": ("gpt-5", "medium")})
    monkeypatch.setattr(memory, "STORE", s.store)
    monkeypatch.setattr(main, "use_postgres_auth", lambda: True)
    monkeypatch.setattr(main, "is_sensitive", lambda q: None)
    monkeypatch.setattr(main, "classify_sensitivity", lambda q, usages=None: s.safety.append(q) or s.flag)
    monkeypatch.setattr(main, "_file_scope_decision", lambda *a, **kw: "denied")

    def scoped_answer(question, msg, client=None):
        s.scoped.append((question, msg["content"]))
        return "คุณอยู่ฝ่ายผลิต", {"model": "gpt-5-mini", "prompt_tokens": 200, "completion_tokens": 10,
                                  "cost_usd": 0.00007}

    monkeypatch.setattr(memory, "answer_from_memory", scoped_answer)
    monkeypatch.setattr(main, "_is_linsip_query", lambda *a: False)
    monkeypatch.setattr(main, "_try_linsip_answer", lambda *a, **kw: None)
    monkeypatch.setattr(main, "_try_url_answer", lambda *a, **kw: None)
    monkeypatch.setattr(main, "_try_web_search_answer", lambda *a, **kw: None)
    monkeypatch.setattr(main, "_agent_enabled", lambda *a: False)
    monkeypatch.setattr(main, "_count_session_turns", lambda sid: 1)
    monkeypatch.setattr(main, "brain_enabled", lambda: False)
    monkeypatch.setattr(main, "search_knowledge", lambda vec, **kw: s.hit)
    monkeypatch.setattr(main, "audit_log", lambda action, **kw: s.audits.append(action))
    monkeypatch.setattr(chat_pg, "ensure_session_pg", lambda **kw: (5, "t"))
    monkeypatch.setattr(chat_pg, "get_session_history_pg", lambda **kw: list(s.history))
    monkeypatch.setattr(chat_pg, "save_chat_message_pg", lambda **kw: s.saved.append(kw) or 77)
    monkeypatch.setattr(chat_pg, "add_usage_to_message_pg", lambda *a: None)
    monkeypatch.setattr(chat_pg, "get_last_image_attachment_pg", lambda sid: None)   # "เพิ่ม…" = image edit?
    monkeypatch.setattr(cm, "history", lambda sid, uid: [{"role": "system", "content": cm.SUMMARY_HEADER + "สรุป"}]
                        + list(s.history))
    monkeypatch.setattr(cm, "schedule_summary", lambda sid, uid: s.summaries.append((sid, uid)))

    def classify(q):
        s.classified.append(q)
        return "general", dict(CLS)

    def fake_stream(kind):
        def run(question, *args, model=None, history=None, usage_holder=None, effort=None, **kw):
            s.streamed.append(NS(kind=kind, model=model, effort=effort, history=history, kw=kw))
            yield "คำ"
            yield "ตอบ"
            usage_holder.update({"model": model, "prompt_tokens": 1000, "completion_tokens": 200, "cost_usd": 0.003,
                                 **({"reasoning_effort": effort} if effort else {})})
        return run

    def fake_answer(kind):
        def run(question, *args, model=None, history=None, effort=None, **kw):
            s.answered.append(NS(kind=kind, model=model, effort=effort, history=history))
            return "คำตอบ", {"model": model, "prompt_tokens": 900, "completion_tokens": 150, "cost_usd": 0.002,
                             **({"reasoning_effort": effort} if effort else {})}
        return run

    monkeypatch.setattr(main, "classify_query", classify)
    monkeypatch.setattr(main, "stream_freely", fake_stream("free"))
    monkeypatch.setattr(main, "stream_from_context", fake_stream("context"))
    monkeypatch.setattr(main, "answer_freely", fake_answer("free"))
    monkeypatch.setattr(main, "answer_from_context", fake_answer("context"))
    return s


def stream(question, user=ADMIN, sid="5"):
    main.app.dependency_overrides[main.current_user] = lambda: user
    try:
        r = TestClient(main.app).post("/chat/stream", data={"message": question, "session_id": sid})
    finally:
        main.app.dependency_overrides.pop(main.current_user, None)
    return [json.loads(line) for line in r.text.splitlines() if line]


def post(question, user=ADMIN):
    main.app.dependency_overrides[main.current_user] = lambda: user
    try:
        return TestClient(main.app).post("/chat", data={"message": question, "session_id": "5"})
    finally:
        main.app.dependency_overrides.pop(main.current_user, None)


# ── V4.1 routing in the pipeline ────────────────────────────────────────────

def test_flags_off_is_the_legacy_pipeline(pipe):
    events = stream("วิธีเบิกค่าเดินทาง")
    assert events[-1]["type"] == "done" and events[-1]["source"] == "llm"
    assert pipe.classified == ["วิธีเบิกค่าเดินทาง"]                    # the classifier still runs
    call = pipe.streamed[0]
    assert call.model == main.LLM_MODEL and call.effort is None         # no effort argument at all
    assert call.history == pipe.history                                 # the plain last-6-turns history
    u = pipe.saved[0]["usage"]
    assert u["route"] == "legacy" and u["route_reason"] == "classifier: general"
    assert u["latency_ms"] >= 0 and u["first_token_ms"] >= 0
    assert u["cost_usd"] == pytest.approx(0.003 + CLS["cost_usd"])
    assert pipe.summaries == []                                         # memory off → no summary


@pytest.mark.parametrize("question,tier,model,effort,source", [
    ("สวัสดีครับ", "light", "gpt-5-mini", "minimal", "llm"),
    ("วิธีเบิกค่าเดินทาง", "standard", "gpt-5", "low", "llm"),
    ("ราคา 1,250 บาท ลด 15% เหลือเท่าไหร่", "calc", "gpt-5-mini", "medium", "llm-calc"),
    ("ช่วยวิเคราะห์ข้อดีข้อเสียของการเพิ่มกะกลางคืน", "deep", "gpt-5", "medium", "llm"),
])
def test_routing_on_picks_model_and_effort_without_the_classifier(pipe, monkeypatch, question, tier,
                                                                  model, effort, source):
    monkeypatch.setattr(mr, "MODEL_ROUTING_MODE", "admin")
    events = stream(question)
    assert events[-1]["source"] == source and "".join(e.get("v", "") for e in events) == "คำตอบ"
    assert pipe.classified == []                                        # no extra AI call
    call = pipe.streamed[0]
    assert (call.model, call.effort) == (model, effort)
    u = pipe.saved[0]["usage"]
    assert u["route"] == tier and u["model_used"] == model and u["reasoning_effort"] == effort
    assert u["cost_usd"] == pytest.approx(0.003)                        # no classifier cost


def test_a_knowledge_hit_is_answered_by_at_least_the_standard_model(pipe, monkeypatch):
    monkeypatch.setattr(mr, "MODEL_ROUTING_MODE", "all")
    pipe.hit = HIT
    events = stream("ขอบคุณครับ")
    assert events[-1]["source"] == "rag"
    call = pipe.streamed[0]
    assert call.kind == "context" and (call.model, call.effort) == ("gpt-5", "low")
    assert "knowledge hit" in pipe.saved[0]["usage"]["route_reason"]


def test_routing_is_per_user(pipe, monkeypatch):
    monkeypatch.setattr(mr, "MODEL_ROUTING_MODE", "admin")
    stream("สวัสดีครับ", user={"id": 9, "role": "user", "username": "staff"})
    assert pipe.classified == ["สวัสดีครับ"] and pipe.streamed[0].effort is None


def test_api_fallback_is_logged(pipe, monkeypatch):
    monkeypatch.setattr(mr, "MODEL_ROUTING_MODE", "admin")

    def rejected(question, *a, model=None, usage_holder=None, effort=None, **kw):
        yield "ตอบ"
        usage_holder.update({"model": "gpt-5", "prompt_tokens": 10, "completion_tokens": 2, "cost_usd": 0.0001,
                             "reasoning_effort": "low", "fallback_from": model})

    monkeypatch.setattr(main, "stream_freely", rejected)
    stream("สวัสดีครับ")
    u = pipe.saved[0]["usage"]
    assert u["model_used"] == "gpt-5" and "API rejected gpt-5-mini → gpt-5" in u["route_reason"]


def test_non_stream_chat_is_routed_too(pipe, monkeypatch):
    monkeypatch.setattr(mr, "MODEL_ROUTING_MODE", "admin")
    r = post("ช่วยวิเคราะห์แผนการผลิตไตรมาสหน้า")
    assert r.status_code == 200 and r.json()["source"] == "llm"
    call = pipe.answered[0]
    assert (call.model, call.effort) == ("gpt-5", "medium") and pipe.classified == []
    assert pipe.saved[0]["usage"]["route"] == "deep"


def test_non_stream_flags_off_keeps_the_legacy_call(pipe):
    r = post("วิธีเบิกค่าเดินทาง")
    assert r.status_code == 200 and pipe.classified == ["วิธีเบิกค่าเดินทาง"]
    assert pipe.answered[0].effort is None and pipe.saved[0]["usage"]["route"] == "legacy"


# ── V4.2 memory & context in the pipeline ───────────────────────────────────

def test_memory_command_is_answered_without_ai(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    events = stream("จำไว้ว่า ฉันอยู่ฝ่ายผลิต กะเช้า")
    assert events[0]["v"].startswith("จำไว้แล้วค่ะ: “ฉันอยู่ฝ่ายผลิต กะเช้า”")
    assert events[-1]["source"] == "memory" and events[-1]["message_id"] == 77
    assert pipe.streamed == [] and pipe.classified == []
    assert pipe.saved[0]["source"] == "memory" and "usage" not in pipe.saved[0]
    assert [i["content"] for i in pipe.store.items(1)] == ["ฉันอยู่ฝ่ายผลิต กะเช้า"]
    assert "memory_command" in pipe.audits


def test_remember_this_takes_the_previous_user_message(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    pipe.history = [{"role": "user", "content": "ผมดูแลเครื่องตัดกระดาษ 3"}, {"role": "assistant", "content": "รับทราบ"}]
    r = post("จำข้อมูลนี้")
    assert r.json()["source"] == "memory"
    assert [i["content"] for i in pipe.store.items(1)] == ["ผมดูแลเครื่องตัดกระดาษ 3"]


def test_memory_off_means_no_commands(pipe):
    events = stream("จำไว้ว่า ฉันอยู่ฝ่ายผลิต")
    assert events[-1]["source"] == "llm" and pipe.streamed and pipe.store.items(1) == []


def test_history_carries_memory_and_summary(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    memory.save(1, "ชอบคำตอบสั้น ๆ")
    memory.save(2, "ข้อมูลของคนอื่น")
    stream("แล้วเรื่องนั้นล่ะ")
    h = pipe.streamed[0].history
    assert h[0]["role"] == "system" and "ชอบคำตอบสั้น ๆ" in h[0]["content"]
    assert all("ข้อมูลของคนอื่น" not in m["content"] for m in h)
    assert h[1]["content"].startswith(cm.SUMMARY_HEADER) and h[2:] == pipe.history
    assert pipe.summaries == [(5, 1)]                                   # summary update scheduled


def test_a_memory_failure_falls_back_to_the_legacy_history(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    monkeypatch.setattr(cm, "history", lambda sid, uid: (_ for _ in ()).throw(RuntimeError("no column")))
    events = stream("คำถาม")
    assert events[-1]["source"] == "llm" and pipe.streamed[0].history == pipe.history
    assert "memory_context_failed" in pipe.audits


def test_agent_turn_gets_the_same_history(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    memory.save(1, "ชอบคำตอบสั้น ๆ")
    monkeypatch.setattr(main, "_agent_tools", lambda u, m: (True, False))
    monkeypatch.setattr(main, "_doc_enabled", lambda u, m: False)
    monkeypatch.setattr(chat_pg, "count_web_searches_today_pg", lambda uid: 0)
    seen = []

    def fake_run(question, history, **kw):
        seen.append(history)
        yield "standard", {"cost_usd": 0.0}

    monkeypatch.setattr(agent, "run_agent", fake_run)
    list(main._agent_turn(ADMIN, "ราคาทองวันนี้", 5, "normal", None))
    assert "ชอบคำตอบสั้น ๆ" in seen[0][0]["content"] and seen[0][1]["content"].startswith(cm.SUMMARY_HEADER)


# ── /memory endpoints ───────────────────────────────────────────────────────

def call(method, path, user=ADMIN, **kw):
    main.app.dependency_overrides[main.current_user] = lambda: user
    try:
        return getattr(TestClient(main.app), method)(path, **kw)
    finally:
        main.app.dependency_overrides.pop(main.current_user, None)


def test_memory_menu_is_hidden_when_off(pipe):
    assert call("get", "/memory").json()["available"] is False
    assert call("post", "/memory", json={"content": "x"}).status_code == 404


def test_memory_endpoints_are_per_user(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    assert call("post", "/memory", json={"content": "ฉันอยู่ฝ่ายผลิต"}).status_code == 200
    r = call("post", "/memory", json={"content": "รหัสผ่านอีเมล abc"})
    assert r.status_code == 400 and "รหัสผ่าน" in r.json()["detail"]
    mine = call("get", "/memory").json()
    assert mine["available"] and mine["enabled"] and [i["content"] for i in mine["items"]] == ["ฉันอยู่ฝ่ายผลิต"]
    mid = mine["items"][0]["id"]
    assert call("get", "/memory", user=OTHER).json()["items"] == []
    assert call("delete", f"/memory/{mid}", user=OTHER).status_code == 404       # not theirs
    assert call("put", "/memory/settings", json={"enabled": False}).json()["enabled"] is False
    assert pipe.store.enabled(1) is False
    assert call("delete", f"/memory/{mid}").status_code == 200
    call("post", "/memory", json={"content": "กะเช้า"})
    assert call("delete", "/memory").json()["deleted"] == 1 and pipe.store.items(1) == []


def test_agent_never_sends_history_together_with_previous_response_id(monkeypatch):
    # The history goes once, in the first call; continuations reference that
    # response and add only the tool output — never the conversation again.
    import url_reader
    from tests.test_agent import FakeClient, resp
    from tests.test_agent import call as tool_call
    from tests.test_agent import stream as answer_stream

    monkeypatch.setattr(url_reader, "fetch_url_text", lambda url: ("T", "text"))
    hist = [{"role": "system", "content": "สรุปบทสนทนา…"}, {"role": "user", "content": "ก่อนหน้า"},
            {"role": "assistant", "content": "ตอบก่อนหน้า"}]
    client = FakeClient(resp(tool_call("read_url", url="https://example.com/a")), answer_stream("คำตอบ"))
    list(agent.run_agent("แล้วหน้านี้ล่ะ", hist, web_quota_left=20, client=client))
    first, second = client.calls
    assert "previous_response_id" not in first
    assert [i["content"] for i in first["input"]] == ["สรุปบทสนทนา…", "ก่อนหน้า", "ตอบก่อนหน้า", "แล้วหน้านี้ล่ะ"]
    assert second["previous_response_id"] == "r0"
    assert all(i["type"] == "function_call_output" for i in second["input"])


# ── usage log: route columns and the routing report ─────────────────────────

def test_route_stats_reach_the_usage_log():
    import usage_log as ul

    sent = []
    cur = NS(execute=lambda sql, args: sent.append((sql, args)))
    ul.insert_usage(cur, user_id=1, feature="chat", source="llm", usage={
        "model_used": "gpt-5-mini", "prompt_tokens": 10, "completion_tokens": 2, "cost_usd": 0.001,
        "route": "light", "route_reason": "greeting/ack", "reasoning_effort": "minimal",
        "latency_ms": 900, "first_token_ms": 400})
    sql, args = sent[0]
    assert "route, route_reason, reasoning_effort, latency_ms, first_token_ms" in sql
    assert args[-5:] == ("light", "greeting/ack", "minimal", 900, 400)
    ul.insert_usage(cur, user_id=1, feature="image", source="image_gen", usage={"model_used": "gpt-image-2"})
    assert sent[1][1][-5:] == (None, None, None, None, None)           # other features: no route


def test_routing_report():
    import usage_log as ul

    rows = [
        {"route": "legacy", "model": "gpt-5", "reasoning_effort": "low", "n": 10, "latency_ms_sum": 80000,
         "latency_n": 10, "first_token_ms_sum": 30000, "first_token_n": 10, "prompt_tokens": 20000,
         "completion_tokens": 9000, "cost_usd": 0.115},
        {"route": "light", "model": "gpt-5-mini", "reasoning_effort": "minimal", "n": 4, "latency_ms_sum": 4000,
         "latency_n": 4, "first_token_ms_sum": None, "first_token_n": 0, "prompt_tokens": 2000,
         "completion_tokens": 300, "cost_usd": 0.0011},
    ]
    out = ul.build_routing(rows, usd_to_thb=36)
    assert [r["route"] for r in out] == ["legacy", "light"]
    legacy, light = out
    assert legacy["avg_latency_ms"] == 8000 and legacy["avg_first_token_ms"] == 3000
    assert legacy["avg_cost_thb"] == pytest.approx(0.414) and light["avg_first_token_ms"] is None
    assert light["label"].startswith("เบา") and light["avg_cost_thb"] == pytest.approx(0.0099)
    report = {**ul.build_report([], since=__import__("datetime").date(2026, 10, 1),
                                until=__import__("datetime").date(2026, 10, 31), seat_usd=30), "routing": out}
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(ul.build_usage_xlsx(report, [])))
    assert "การเลือกโมเดล" in wb.sheetnames
    assert wb["การเลือกโมเดล"]["A2"].value == "เดิม (ตัวจัดหมวด AI)"


def test_context_summary_cost_is_not_a_request():
    import usage_log as ul

    assert not ul._is_request("context", "summary")
    assert ul._is_request("chat", "llm") and not ul._is_request("safety", "classifier")


# ── memory and the safety gate ──────────────────────────────────────────────

def test_memory_command_runs_before_the_ai_classifier(pipe, monkeypatch):
    # The AI classifier blocked "จำไว้ว่า ฉันอยู่ฝ่ายผลิต…" as personal data in the
    # real E2E; the command is the user's own data and needs no AI.
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    pipe.flag = ("PII", "personal info")
    events = stream("จำไว้ว่า ฉันอยู่ฝ่ายผลิต ทำงานกะเช้า")
    assert events[-1]["source"] == "memory" and pipe.safety == []
    assert [i["content"] for i in pipe.store.items(1)] == ["ฉันอยู่ฝ่ายผลิต ทำงานกะเช้า"]


def test_keyword_filter_still_comes_first(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    monkeypatch.setattr(main, "is_sensitive", lambda q: "เงินเดือน")
    events = stream("จำไว้ว่า เงินเดือนของพี่สมชาย 50,000")
    assert events[-1]["source"] == "blocked" and pipe.store.items(1) == []


def test_flagged_question_about_myself_is_answered_from_my_memory_only(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    memory.save(1, "ฉันอยู่ฝ่ายผลิต ทำงานกะเช้า")
    memory.save(2, "ข้อมูลของคนอื่น")
    pipe.flag = ("PII", "personal info")
    events = stream("ฉันอยู่ฝ่ายไหน")
    assert events[-1]["source"] == "memory_scoped" and events[0]["v"] == "คุณอยู่ฝ่ายผลิต"
    question, given = pipe.scoped[0]
    assert question == "ฉันอยู่ฝ่ายไหน" and "ฝ่ายผลิต" in given and "ข้อมูลของคนอื่น" not in given
    assert pipe.streamed == [] and pipe.saved[0]["source"] == "memory_scoped"
    assert pipe.saved[0]["usage"]["model_used"] == "gpt-5-mini"


@pytest.mark.parametrize("question,mode,saved", [
    ("เงินเดือนของพี่สมชายเท่าไหร่", "admin", True),       # not about the user → blocked as before
    ("ฉันอยู่ฝ่ายไหน", "off", True),                      # memory off → blocked as before
    ("ฉันอยู่ฝ่ายไหน", "admin", False),                   # nothing remembered → blocked as before
])
def test_other_flagged_questions_stay_blocked(pipe, monkeypatch, question, mode, saved):
    monkeypatch.setattr(memory, "MEMORY_MODE", mode)
    if saved:
        pipe.store.add(1, "ฉันอยู่ฝ่ายผลิต", None)
        pipe.store.set_enabled(1, True)
    pipe.flag = ("PII", "personal info")
    events = stream(question)
    assert events[-1]["source"] == "blocked" and events[0]["v"] == main.BLOCKED_RESPONSE
    assert pipe.scoped == []


def test_non_stream_flagged_question_about_myself(pipe, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_MODE", "admin")
    memory.save(1, "ฉันอยู่ฝ่ายผลิต")
    pipe.flag = ("PII", "personal info")
    r = post("ฉันอยู่ฝ่ายไหน")
    assert r.json()["source"] == "memory_scoped" and r.json()["answer"] == "คุณอยู่ฝ่ายผลิต"
    r = post("จำไว้ว่า ฉันทำงานกะเช้า")
    assert r.json()["source"] == "memory" and len(pipe.store.items(1)) == 2

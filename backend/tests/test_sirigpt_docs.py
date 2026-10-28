"""
Tests for sirigpt_docs — the document helpers the sandbox builds files with and
the server checks them with. One section per file type (Word → Excel →
PowerPoint → PDF), then edits that must not lose content. Offline.
"""

import io
import zipfile

import pytest
from PIL import Image

import sirigpt_docs as sd


@pytest.fixture
def png(tmp_path):
    p = tmp_path / "chart.png"
    Image.new("RGB", (800, 450), (91, 44, 159)).save(p)
    return str(p)


# ═════════════════════════════════════════════════════════════════════ Word ══

def _thai_report(path, png):
    doc = sd.new_docx()
    sd.heading(doc, "รายงานยอดขายไตรมาส 1/2569", 0)
    sd.heading(doc, "1. สรุปผล", 1)
    sd.para(doc, "ยอดขายรวมทั้งไตรมาสเท่ากับ **97,650 บาท** เพิ่มขึ้นจากไตรมาสก่อน")
    sd.bullets(doc, ["ม.ค. สูงที่สุด", "มี.ค. ต่ำที่สุด"])
    sd.table(doc, [["เดือน", "ยอดขาย (บาท)"], ["ม.ค.", 41650], ["ก.พ.", 33400], ["มี.ค.", 22600.5]])
    sd.picture(doc, png, width_cm=14, caption="ภาพที่ 1 ยอดขายรายเดือน")
    sd.page_break(doc)
    sd.heading(doc, "2. ข้อเสนอแนะ", 1)
    sd.para(doc, "เพิ่มกำลังการผลิตงานรวมเล่ม", align="justify")
    doc.save(path)
    return path


def test_word_report_is_thai_ready_and_complete(tmp_path, png):
    path = _thai_report(str(tmp_path / "รายงาน.docx"), png)
    r = sd.check(path)
    assert r["ok"], r
    s = r["stats"]
    assert s["tables"] == 1 and s["images"] == 1 and s["page_numbers"] and s["thai"]
    assert s["headings"] >= 3
    assert not any("ฟอนต์" in w for w in r["warnings"])            # Thai (cs) font is set

    with zipfile.ZipFile(path) as z:
        styles, body = z.read("word/styles.xml").decode(), z.read("word/document.xml").decode()
        footer = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
    assert 'w:cs="Sarabun"' in styles and "asciiTheme" not in styles.split('w:styleId="Heading1"')[1][:600]
    assert "PAGE" in footer and "NUMPAGES" in footer
    assert "w:tblHeader" in body                                     # header row repeats


def test_word_table_formats_numbers(tmp_path, png):
    from docx import Document

    path = _thai_report(str(tmp_path / "r.docx"), png)
    t = Document(path).tables[0]
    assert t.cell(1, 1).text == "41,650" and t.cell(3, 1).text == "22,600.50"
    assert t.cell(1, 1).paragraphs[0].alignment == 2                  # right


def test_word_run_properties_stay_in_schema_order(tmp_path, png):
    """Word refuses a document whose w:rPr children are out of order."""
    import re

    path = _thai_report(str(tmp_path / "r.docx"), png)
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode() + z.read("word/styles.xml").decode()
    rank = {t: i for i, t in enumerate(sd._RPR_ORDER)}
    for rpr in re.findall(r"<w:rPr>(.*?)</w:rPr>", xml):
        tags = [t for t in re.findall(r"<w:(\w+)", rpr) if t in rank]
        assert tags == sorted(tags, key=rank.get), tags


def test_word_with_user_template_keeps_its_setup(tmp_path):
    from docx import Document
    from docx.shared import Cm

    tpl = Document()
    tpl.sections[0].left_margin = Cm(3.5)
    tpl.add_paragraph("หัวจดหมายบริษัท")
    tpl.save(tmp_path / "template.docx")
    doc = sd.new_docx(template=str(tmp_path / "template.docx"))
    sd.para(doc, "เนื้อหาใหม่")
    doc.save(tmp_path / "out.docx")
    out = Document(str(tmp_path / "out.docx"))
    assert abs(out.sections[0].left_margin - Cm(3.5)) < Cm(0.01)        # Word stores twips
    assert [p.text for p in out.paragraphs][:2] == ["หัวจดหมายบริษัท", "เนื้อหาใหม่"]


def test_word_without_thai_font_is_flagged(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("ภาษาไทยแต่ไม่ได้ตั้งฟอนต์")
    d.save(tmp_path / "plain.docx")
    r = sd.check(str(tmp_path / "plain.docx"))
    assert r["ok"] and any("ฟอนต์ภาษาไทย" in w for w in r["warnings"])


def test_broken_and_empty_files_are_not_ok(tmp_path):
    (tmp_path / "bad.docx").write_bytes(b"not a zip")
    assert sd.check(str(tmp_path / "bad.docx"))["ok"] is False
    from docx import Document

    Document().save(tmp_path / "empty.docx")
    r = sd.check(str(tmp_path / "empty.docx"))
    assert r["ok"] is False and "ไม่มีเนื้อหา" in r["issues"][0]


# ═══════════════════════════════════════════════════════════════════ Excel ══

def _sales_workbook(path):
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    data = {"ม.ค.": [("กล่อง", 10, 1200), ("ฉลาก", 25, 450)], "ก.พ.": [("กล่อง", 7, 1200), ("ฉลาก", 30, 450)]}
    for name, rows in data.items():
        ws = wb.create_sheet(name)
        ws.append(["สินค้า", "จำนวน", "ราคา", "ยอดขาย"])
        for i, (item, qty, price) in enumerate(rows, start=2):
            ws.append([item, qty, price, f"=B{i}*C{i}"])
        ws.append(["รวม", "=SUM(B2:B3)", None, "=SUM(D2:D3)"])
        sd.style_table(ws)
    s = wb.create_sheet("สรุป")
    s.append(["เดือน", "ยอดขาย", "สัดส่วน"])
    s.append(["ม.ค.", "='ม.ค.'!D4", "=ROUND(B2/B4*100,2)"])
    s.append(["ก.พ.", "=ก.พ.!D4", "=ROUND(B3/B4*100,2)"])
    s.append(["รวม", "=SUM(B2:B3)", "=IF(B4>0,100,0)"])
    s["E1"], s["E2"] = "ยอดกล่อง", "=SUMIF('ม.ค.'!A2:A3,\"กล่อง\",'ม.ค.'!D2:D3)+SUMIF('ก.พ.'!A2:A3,\"กล่อง\",'ก.พ.'!D2:D3)"
    ch = BarChart()
    ch.add_data(Reference(s, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=3))
    s.add_chart(ch, "G2")
    sd.prepare_workbook(wb)
    wb.save(path)
    return path


def test_excel_formulas_are_evaluated_correctly(tmp_path):
    path = _sales_workbook(str(tmp_path / "ยอดขาย.xlsx"))
    rep = sd.formula_report(path)
    v = rep["values"]
    assert rep["errors"] == [] and rep["unverified"] == []
    assert rep["formulas"] == rep["verified"] == 15  # 4 per month sheet x2 + 7 on สรุป
    assert v["ม.ค.!D4"] == 23250 and v["ก.พ.!D4"] == 21900          # 12000+11250, 8400+13500
    assert v["สรุป!B4"] == 45150 and v["สรุป!C2"] == 51.5 and v["สรุป!C3"] == 48.5
    assert v["สรุป!E2"] == 20400                                       # SUMIF across sheets


def test_excel_check_reports_charts_and_formulas(tmp_path):
    path = _sales_workbook(str(tmp_path / "b.xlsx"))
    r = sd.check(path)
    assert r["ok"], r
    assert r["stats"]["sheets"] == ["ม.ค.", "ก.พ.", "สรุป"] and r["stats"]["charts"] == 1
    assert r["stats"]["formulas"] == r["stats"]["formulas_verified"] == 15
    import openpyxl

    assert openpyxl.load_workbook(path).calculation.fullCalcOnLoad is True


def test_excel_formula_errors_are_issues(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([1, 0, "=A1/B1", "=XLOOKUP(A1,A1:A2,B1:B2)", "=C1+1"])
    wb.save(tmp_path / "e.xlsx")
    r = sd.check(str(tmp_path / "e.xlsx"))
    assert r["ok"] is False
    assert any("#DIV/0!" in i for i in r["issues"])
    assert any("ตรวจผลสูตรไม่ได้ 1" in w for w in r["warnings"])       # XLOOKUP not supported


@pytest.mark.parametrize("formula,expected", [
    ("=1+2*3", 7), ("=(1+2)*3", 9), ("=2^3", 8), ("=-A1+10", 5), ("=50%", 0.5),
    ("=A1&\"บาท\"", "5บาท"), ("=IF(A1>3,\"สูง\",\"ต่ำ\")", "สูง"), ("=IFERROR(1/0,\"-\")", "-"),
    ("=IF(A1>3,1,1/0)", 1), ("=ROUND(2.5,0)", 3), ("=ROUND(-2.5,0)", -3), ("=ROUNDDOWN(2.99,1)", 2.9),
    ("=AVERAGE(A1:A3)", 5), ("=COUNT(A1:A3)", 1), ("=COUNTA(A1:A3)", 1), ("=MAX(A1:A3)+MIN(A1:A3)", 10),
    ("=COUNTIF(A1:A3,\">4\")", 1), ("=ABS(-3)", 3), ("=TRUE", True),
    # an error inside ( ) in the branch not taken (prod: =IF(ROW()=2,"-",IF(B1=0,"",(B2-B1)/B1)), B1 a header)
    ("=IF(A1>3,1,(A1-\"x\")/2)", 1), ("=IF(A1>3,\"-\",IF(A1=0,\"\",(A1-\"x\")/A1))", "-"),
])
def test_formula_evaluator(tmp_path, formula, expected):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = 5, formula
    wb.save(tmp_path / "f.xlsx")
    rep = sd.formula_report(str(tmp_path / "f.xlsx"))
    assert rep["errors"] == [] and rep["unverified"] == [], rep
    assert rep["values"]["Sheet!B1"] == expected


def test_lookup_and_conditional_functions(tmp_path):
    # A prod dashboard used ROW / INDEX+MATCH and was reported as "ตรวจผลสูตรไม่ได้ 7 ช่อง".
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ข้อมูล"
    for row in [["เดือน", "สาขา", "ยอดขาย"], ["ส.ค.", "กทม", 1250000], ["ก.ย.", "กทม", 1420000],
                ["ต.ค.", "ชลบุรี", 1580000]]:
        ws.append(row)
    cases = {
        "E1": ('=IF(ROW()=1,"-",C1-C0)', "-"), "E3": ("=IF(ROW()=2,\"-\",C3-C2)", 170000),
        "F1": ("=COLUMN()", 6),
        "E5": ("=INDEX(A2:A4,MATCH(MAX(C2:C4),C2:C4,0))", "ต.ค."),
        "E6": ("=INDEX(A1:C4,3,3)", 1420000), "E7": ("=MATCH(1500000,C2:C4,1)", 2),
        "E8": ('=VLOOKUP("ก.ย.",A2:C4,3,FALSE)', 1420000), "E9": ("=VLOOKUP(\"ก.ย.\",A2:C4,2,0)", "กทม"),
        "E10": ('=SUMIFS(C2:C4,B2:B4,"กทม")', 2670000), "E11": ('=COUNTIFS(B2:B4,"กทม",C2:C4,">1300000")', 1),
        "E12": ('=AVERAGEIFS(C2:C4,B2:B4,"กทม")', 1335000), "E13": ("=MEDIAN(C2:C4)", 1420000),
        "E14": ("=AND(C2>0,C3>C2)", True), "E15": ("=OR(C2>2000000,NOT(C3>0))", False),
        "E16": ('=LEN(B4)&"-"&CONCAT(A2,"/",B2)', "6-ส.ค./กทม"),
    }
    for cell, (f, _) in cases.items():
        ws[cell] = f
    ws["E2"] = '=MATCH("ธ.ค.",A2:A4,0)'                                 # not found → #N/A
    wb.save(tmp_path / "l.xlsx")
    rep = sd.formula_report(str(tmp_path / "l.xlsx"))
    assert rep["errors"] == [{"cell": "ข้อมูล!E2", "error": "#N/A"}]
    assert [u["cell"] for u in rep["unverified"]] == ["ข้อมูล!E1"]       # C0 is not a cell
    for cell, (_, want) in cases.items():
        if cell != "E1":
            assert rep["values"][f"ข้อมูล!{cell}"] == want, (cell, rep["values"][f"ข้อมูล!{cell}"])


def test_style_table_keeps_percent_and_fits_numbers(tmp_path):
    # Seen on prod: a "(%)" column of ratios was shown as 0.14, and a 10.5-wide
    # column held 1,250,000.00 (Excel shows ####).
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["เดือน", "ยอดขาย (บาท)", "เปลี่ยนแปลง (%)", "สัดส่วน (%)", "ร้อยละ (คิดแล้ว)"])
    ws.append(["ส.ค.", 1250000, "-", 0.3, 30])
    ws.append(["ก.ย.", 1420000, "=(B3-B2)/B2", 0.7, "=B3/B2*100"])
    ws["B3"].number_format = "#,##0"                                     # set by the model: kept
    sd.style_table(ws)
    assert ws["C3"].number_format == "0.00%" and ws["D2"].number_format == "0.00%"
    assert ws["E2"].number_format == "#,##0.00" and ws["E3"].number_format == "#,##0.00"   # already ×100
    assert ws["B2"].number_format == "#,##0.00" and ws["B3"].number_format == "#,##0"
    assert ws.column_dimensions["B"].width >= len("1,250,000.00") * 1.1
    assert ws.column_dimensions["C"].width < 20                          # not sized by the formula text


def test_circular_reference_is_unverified_not_a_crash(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = "=A2+1", "=A1+1"
    wb.save(tmp_path / "c.xlsx")
    rep = sd.formula_report(str(tmp_path / "c.xlsx"))
    assert rep["verified"] == 0 and len(rep["unverified"]) == 2


# ═══════════════════════════════════════════════════════════════ PowerPoint ══

def _deck(path, png):
    prs = sd.new_pptx()
    sd.title_slide(prs, "สรุปผลการดำเนินงานไตรมาส 1", "ฝ่ายขาย · บริษัท ศิริวัฒนาอินเตอร์พริ้นท์")
    sd.bullet_slide(prs, "ประเด็นสำคัญ", ["ยอดขายรวม 97,650 บาท", "เดือนมกราคมสูงสุด 41,650 บาท",
                                          ("กล่องลูกฟูกเติบโตดี", 1), "ต้องเร่งงานรวมเล่มในไตรมาสถัดไป"])
    sd.table_slide(prs, "ยอดขายรายเดือน", [["เดือน", "ยอดขาย"], ["ม.ค.", 41650], ["ก.พ.", 33400], ["มี.ค.", 22600]])
    sd.chart_slide(prs, "แนวโน้มยอดขาย", ["ม.ค.", "ก.พ.", "มี.ค."], {"ยอดขาย": [41650, 33400, 22600]})
    sd.image_slide(prs, "ภาพประกอบ", png, caption="ภาพจากโรงงาน")
    prs.save(path)
    return path


def test_five_slide_deck_is_editable_and_fits(tmp_path, png):
    path = _deck(str(tmp_path / "สไลด์.pptx"), png)
    r = sd.check(path)
    assert r["ok"], r
    s = r["stats"]
    assert s["slides"] == 5 and s["tables"] == 1 and s["charts"] == 1 and s["pictures"] == 1
    assert s["text_boxes"] >= 6                                       # titles + bullets are text, not images
    assert r["warnings"] == []                                        # nothing overflows or overlaps


def test_word_bullets_and_headings_have_no_literal_markdown(tmp_path):
    # Seen on prod: "**ยอดขายรวม**: …" bullets kept their asterisks.
    from docx import Document

    doc = sd.new_docx()
    sd.heading(doc, "**บทสรุป**", 1)
    sd.bullets(doc, ["**ยอดขายรวม**: 97,650 บาท", "ข้อสอง"])
    doc.save(tmp_path / "b.docx")
    d = Document(str(tmp_path / "b.docx"))
    assert "**" not in "\n".join(p.text for p in d.paragraphs)
    first = next(p for p in d.paragraphs if p.style.name == "List Bullet")
    assert first.text == "ยอดขายรวม: 97,650 บาท" and first.runs[0].text == "ยอดขายรวม" and first.runs[0].bold


def test_kpi_cards_and_bold_bullets(tmp_path):
    from pptx import Presentation

    prs = sd.new_pptx()
    sd.kpi_slide(prs, "ตัวเลขสำคัญ", [("ยอดขายรวม Q1", 97650), ("สูงสุด (ม.ค.)", 41650), ("ต่ำสุด (มี.ค.)", 22600)],
                 points=["**ยอดขายลดลง** ต่อเนื่องทุกเดือน", "ต้องเร่งแผนไตรมาส 2"])
    sd.kpi_slide(prs, "สรุป", [("ยอดขาย", 97650.5)])
    sd.bullet_slide(prs, "**แผนไตรมาส 2**", ["**เป้าหมาย:** เพิ่มยอดขาย 10%"])
    prs.save(tmp_path / "k.pptx")
    r = sd.check(str(tmp_path / "k.pptx"))
    assert r["ok"] and r["warnings"] == [], r["warnings"]
    deck = Presentation(str(tmp_path / "k.pptx"))
    text = "\n".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert "**" not in text                                           # seen literally on an AI-made slide
    assert "ยอดขายรวม Q1\n97,650" in text and "97,650.50" in text
    runs = [r for sh in deck.slides[2].shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs]
    assert any(r.text == "เป้าหมาย:" and r.font.bold for r in runs)


def test_boxes_placed_on_top_of_slide_content_are_flagged(tmp_path):
    # Seen in the E2E deck: number boxes drawn over the bullets, a note over the chart.
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = sd.new_pptx()
    s1 = sd.bullet_slide(prs, "ประเด็นสำคัญ", ["ยอดขายรวม 97,650 บาท", "มกราคมสูงสุด"])
    s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(3), Inches(1)).text_frame.text = "97,650"
    s2 = sd.chart_slide(prs, "แนวโน้ม", ["ม.ค.", "ก.พ."], {"ยอดขาย": [41650, 33400]})
    s2.shapes.add_textbox(Inches(1), Inches(5), Inches(5), Inches(0.5)).text_frame.text = "ยอดขายลดลงต่อเนื่อง"
    s3 = sd.bullet_slide(prs, "ไม่ซ้อน", ["ข้อความสั้น"])     # a big text box with one line: the rest is free
    s3.shapes.add_textbox(Inches(1), Inches(5.5), Inches(5), Inches(0.5)).text_frame.text = "หมายเหตุท้ายสไลด์"
    prs.save(tmp_path / "o.pptx")
    w = sd.check(str(tmp_path / "o.pptx"))["warnings"]
    assert "สไลด์ 1: ข้อความ/รูป/ตาราง/กราฟวางซ้อนทับกัน" in w
    assert "สไลด์ 2: ข้อความ/รูป/ตาราง/กราฟวางซ้อนทับกัน" in w
    assert not any(x.startswith("สไลด์ 3") for x in w), w


def test_slide_text_uses_thai_font_slot(tmp_path, png):
    path = _deck(str(tmp_path / "d.pptx"), png)
    with zipfile.ZipFile(path) as z:
        xml = z.read("ppt/slides/slide2.xml").decode()
    assert f'<a:cs typeface="{sd.PPTX_FONT}"/>' in xml


def test_long_bullets_shrink_to_fit():
    long = ["ข้อความยาวมากเพื่อทดสอบการตัดบรรทัดของภาษาไทยในสไลด์ " * 3] * 6
    size = sd.fit_font_size(long, 11.5 * 72, 4.8 * 72, max_size=28, min_size=12)
    assert 12 <= size < 28
    assert sd.estimate_lines(long, size, 11.5 * 72) * size * 1.25 <= 4.8 * 72 + 6 * size


def test_overflowing_and_image_only_slides_are_flagged(tmp_path, png):
    from pptx.util import Inches, Pt

    prs = sd.new_pptx()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = "ข้อความยาวเกินกรอบมาก " * 30
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
    prs.save(tmp_path / "bad.pptx")
    r = sd.check(str(tmp_path / "bad.pptx"))
    assert any("สไลด์ 1: ข้อความอาจล้นกรอบ" == w for w in r["warnings"])
    assert any("สไลด์ 2: เป็นภาพทั้งหน้า" in i for i in r["issues"]) and r["ok"] is False


# ═════════════════════════════════════════════════════════════════════ PDF ══

def test_pdf_pages_text_and_scans(tmp_path, png):
    import fitz

    doc = fitz.open()
    for i in range(3):
        doc.new_page().insert_text((72, 72), f"Page {i + 1} report")
    doc.save(tmp_path / "a.pdf")
    r = sd.check(str(tmp_path / "a.pdf"))
    assert r["ok"] and r["stats"]["pages"] == 3 and r["stats"]["text_chars"] > 20

    scan = fitz.open()
    page = scan.new_page()
    page.insert_image(page.rect, filename=png)
    scan.save(tmp_path / "scan.pdf")
    r = sd.check(str(tmp_path / "scan.pdf"))
    assert r["ok"] and any("ภาพสแกน" in w for w in r["warnings"])


# ═════════════════════════════════════════════ edits keep the original content ══

def test_word_edit_in_place_keeps_everything(tmp_path, png):
    from docx import Document

    orig = _thai_report(str(tmp_path / "orig.docx"), png)
    d = Document(orig)
    d.paragraphs[2].runs[0].text = "ยอดขายรวม (แก้ไข) "
    sd.para(d, "ภาคผนวก: เพิ่มเติม")
    d.save(tmp_path / "orig_แก้ไข.docx")
    r = sd.check(str(tmp_path / "orig_แก้ไข.docx"), original=orig)
    assert r["ok"] and not [w for w in r["warnings"] if "ลดลง" in w or "หาย" in w], r["warnings"]


def test_word_rebuilt_from_scratch_is_reported(tmp_path, png):
    orig = _thai_report(str(tmp_path / "orig.docx"), png)
    doc = sd.new_docx()
    sd.para(doc, "สรุปสั้น ๆ")
    doc.save(tmp_path / "new.docx")
    w = sd.check(str(tmp_path / "new.docx"), original=orig)["warnings"]
    assert "ตาราง ลดลงจาก 1 เหลือ 0" in w and "รูปภาพ ลดลงจาก 1 เหลือ 0" in w


def test_excel_edit_keeps_formulas_charts_and_sheets(tmp_path):
    import openpyxl

    orig = _sales_workbook(str(tmp_path / "orig.xlsx"))
    wb = openpyxl.load_workbook(orig)
    wb["ม.ค."]["B2"] = 12                                             # a normal edit
    wb.save(tmp_path / "edited.xlsx")
    r = sd.check(str(tmp_path / "edited.xlsx"), original=orig)
    assert r["ok"] and r["warnings"] == [], r["warnings"]
    assert sd.formula_report(str(tmp_path / "edited.xlsx"))["values"]["ม.ค.!D4"] == 25650

    wb = openpyxl.load_workbook(orig)
    del wb["ก.พ."]
    wb["สรุป"]._charts = []
    wb["สรุป"]["B2"] = 23250                                           # formula replaced by a number
    wb.save(tmp_path / "lossy.xlsx")
    w = sd.check(str(tmp_path / "lossy.xlsx"), original=orig)["warnings"]
    assert "ชีตหายไป: ก.พ." in w and "กราฟ ลดลงจาก 1 เหลือ 0" in w
    assert any(x.startswith("สูตร ลดลง") for x in w)


def test_pptx_edit_keeps_slides_and_flags_a_deleted_one(tmp_path, png):
    from pptx import Presentation

    orig = _deck(str(tmp_path / "orig.pptx"), png)
    prs = Presentation(orig)
    prs.slides[1].shapes[0].text_frame.text = "ประเด็นสำคัญ (ปรับปรุง)"
    prs.save(tmp_path / "ok.pptx")
    assert sd.check(str(tmp_path / "ok.pptx"), original=orig)["warnings"] == []

    prs = Presentation(orig)
    sld = prs.slides._sldIdLst
    sld.remove(sld[3])                                                 # drop the chart slide
    prs.save(tmp_path / "less.pptx")
    w = sd.check(str(tmp_path / "less.pptx"), original=orig)["warnings"]
    assert "สไลด์ ลดลงจาก 5 เหลือ 4" in w and "กราฟ ลดลงจาก 1 เหลือ 0" in w

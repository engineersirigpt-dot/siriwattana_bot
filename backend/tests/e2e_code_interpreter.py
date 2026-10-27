"""
End-to-end check of Code Interpreter V2 against a RUNNING backend (real OpenAI
calls, costs a few US cents per run). Not collected by pytest.

Expects a throwaway postgres deployment started with
    CODE_INTERPRETER_MODE=admin  CODE_INTERPRETER_USERS=e2e_trial  AGENT_V1_MODE=admin
and three users: e2e_admin (role admin), e2e_trial, e2e_other (both role user).

    python tests/e2e_code_interpreter.py --base http://host:8000 --setup   # register users
    python tests/e2e_code_interpreter.py --base http://host:8000           # flags on
    python tests/e2e_code_interpreter.py --base http://host:8000 --phase off  # after restarting
                                                                         # with both flags off
Prints one JSON line per case and a PASS/FAIL summary (exit 1 on any failure).
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx
import openpyxl
from PIL import Image, ImageDraw, ImageFont

PASSWORD = "e2e-pass-1234"
USERS = ("e2e_admin", "e2e_trial", "e2e_other")
FONT = Path(__file__).resolve().parent.parent / "fonts" / "Sarabun-Regular.ttf"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BASE = ""
results: list[dict] = []


# ── HTTP helpers ─────────────────────────────────────────────────────────────

def login(user: str) -> str:
    r = httpx.post(f"{BASE}/auth/login", data={"username": user, "password": PASSWORD}, timeout=30)
    r.raise_for_status()
    return r.json()["access_token"]


def hdr(tok: str) -> dict:
    return {"Authorization": f"Bearer {tok}"}


def stream_chat(tok: str, msg: str, sid: int | None = None) -> dict:
    t0 = time.monotonic()
    first = None
    statuses, parts, done = [], [], None
    data = {"message": msg}
    if sid:
        data["session_id"] = str(sid)
    with httpx.stream("POST", f"{BASE}/chat/stream", data=data, headers=hdr(tok), timeout=400) as r:
        r.raise_for_status()
        for line in r.iter_lines():
            if not line:
                continue
            ev = json.loads(line)
            if first is None:
                first = time.monotonic() - t0
            if ev["type"] == "status":
                statuses.append(ev["state"])
            elif ev["type"] == "delta":
                parts.append(ev["v"])
            elif ev["type"] == "done":
                done = ev
            elif ev["type"] == "error":
                raise RuntimeError(f"stream error: {ev}")
    return {
        "answer": "".join(parts), "source": (done or {}).get("source"), "statuses": statuses,
        "attachments": (done or {}).get("attachments") or [], "session_id": (done or {}).get("session_id"),
        "message_id": (done or {}).get("message_id"),
        "first_s": round(first or 0, 2), "total_s": round(time.monotonic() - t0, 2),
    }


def upload_chat(tok: str, msg: str, files: list[tuple[str, bytes, str]], sid: int | None = None) -> dict:
    t0 = time.monotonic()
    data = {"message": msg}
    if sid:
        data["session_id"] = str(sid)
    r = httpx.post(
        f"{BASE}/chat", data=data, headers=hdr(tok), timeout=400,
        files=[("files", f) for f in files],
    )
    r.raise_for_status()
    j = r.json()
    return {
        "answer": j["answer"], "source": j["source"], "attachments": j.get("attachments") or [],
        "session_id": j["session_id"], "message_id": j.get("message_id"),
        "statuses": [], "first_s": None, "total_s": round(time.monotonic() - t0, 2),
    }


def download(tok: str, aid: int) -> httpx.Response:
    return httpx.get(f"{BASE}/attachments/{aid}", headers=hdr(tok), timeout=60)


def digits(text: str) -> str:
    return text.replace(",", "").replace(" ", "")


def record(case: str, ok: bool, res: dict | None = None, **extra) -> bool:
    row = {"case": case, "ok": bool(ok)}
    if res:
        row.update({
            "source": res.get("source"), "message_id": res.get("message_id"),
            "first_s": res.get("first_s"), "total_s": res.get("total_s"),
            "statuses": res.get("statuses"),
            "files": [a["filename"] for a in res.get("attachments", []) if a.get("generated")],
        })
    row.update(extra)
    if not ok and res:
        row["answer"] = res.get("answer", "")[:600]
    results.append(row)
    print(json.dumps(row, ensure_ascii=False), flush=True)
    return ok


ONLY: list[str] = []   # --only: run just the cases whose name contains one of these


def guarded(case: str, fn) -> None:
    if ONLY and not any(o in case for o in ONLY):
        return
    try:
        fn()
    except Exception as e:  # a crash is a failed case, not an aborted run
        record(case, False, error=f"{type(e).__name__}: {e}"[:500])


# ── Test data ────────────────────────────────────────────────────────────────

SALES = {  # sheet → rows (product, qty, unit price); each sheet ends with a "รวม" row
    "ม.ค.": [("กล่องลูกฟูก", 10, 1200), ("ฉลากสินค้า", 25, 450), ("แคตตาล็อก", 8, 2300)],   # 41,650
    "ก.พ.": [("กล่องลูกฟูก", 7, 1200), ("ฉลากสินค้า", 30, 450), ("แคตตาล็อก", 5, 2300)],    # 33,400
    "มี.ค.": [("กล่องลูกฟูก", 9, 1200), ("ฉลากสินค้า", 16, 450), ("แคตตาล็อก", 2, 2300)],   # 22,600
}
SALES_TOTAL = 97_650

DEFECTS = [("สีเพี้ยน", 120), ("ตัดตกขอบ", 85), ("กาวไม่ติด", 60), ("รอยขีดข่วน", 40),
           ("พิมพ์เบลอ", 25), ("กระดาษยับ", 15), ("อื่นๆ", 5)]   # total 350, cum% 34.29 58.57 75.71 87.14 …


def sales_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, rows in SALES.items():
        ws = wb.create_sheet(sheet)
        ws.append(["สินค้า", "จำนวน", "ราคาต่อหน่วย", "ยอดขาย"])
        for name, qty, price in rows:
            ws.append([name, qty, price, qty * price])
        ws.append(["รวม", None, None, sum(q * p for _, q, p in rows)])   # a total row that must not be double-counted
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def defects_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["ประเภทของเสีย", "จำนวน"])
    w.writerows(DEFECTS)
    return buf.getvalue().encode("utf-8-sig")


def typo_png() -> bytes:
    img = Image.new("RGB", (1400, 260), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(str(FONT), 54)
    d.text((40, 40), "ขอแสดงความยินดีกับพนักงานดีเด่นประจำปี", fill="black", font=font)
    d.text((40, 140), "ทุกท่านที่ได้รับรางวันในครั้งนี้", fill="black", font=font)   # รางวัล → รางวัน
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


# ── Phase "on" ───────────────────────────────────────────────────────────────

CALCS = [  # (case, question, expected digits)
    ("calc_compound_interest",
     "ฝากเงิน 100,000 บาท ดอกเบี้ย 5% ต่อปี ทบต้นทุกปี 10 ปี จะได้เงินรวมเท่าไหร่ (ทศนิยม 2 ตำแหน่ง)",
     ["162889.46"]),
    ("calc_loan_payment",
     "กู้เงิน 3,000,000 บาท ดอกเบี้ย 6% ต่อปี ผ่อนรายเดือน 30 ปี ค่างวดต่อเดือนเท่าไหร่ (ทศนิยม 2 ตำแหน่ง)",
     ["17986.52"]),
    ("calc_sample_stdev",
     "หาค่าเฉลี่ยและส่วนเบี่ยงเบนมาตรฐานแบบตัวอย่าง (sample SD) ของ 12, 15, 18, 22, 30 ทศนิยม 2 ตำแหน่ง",
     ["19.4", "6.99"]),
]

GENERAL = [
    "สวัสดีครับ",
    "ช่วยเขียนอีเมลขอลาพักร้อน 2 วันถึงหัวหน้าให้หน่อย",
    "Python คืออะไร อธิบายสั้นๆ 2 บรรทัด",
    "12 + 30 เท่ากับเท่าไหร่",
    "แปลเป็นภาษาอังกฤษ: ขอบคุณที่ช่วยเหลือครับ",
]


def phase_on(tok: dict[str, str]) -> None:
    admin, trial, other = tok["e2e_admin"], tok["e2e_trial"], tok["e2e_other"]

    # 1) Known-answer calculations (streaming, admin: Agent V1 + CI both on)
    for case, q, expect in CALCS:
        def run_calc(case=case, q=q, expect=expect):
            res = stream_chat(admin, q)
            record(case, all(e in digits(res["answer"]) for e in expect), res)
        guarded(case, run_calc)

    # 2) General questions must not start the sandbox
    for i, q in enumerate(GENERAL):
        def run_general(i=i, q=q):
            res = stream_chat(admin, q)
            record(f"general_no_ci_{i}", res["source"] != "code" and "agent_coding" not in res["statuses"], res)
        guarded(f"general_no_ci_{i}", run_general)

    # 3) Multi-sheet Excel totals (CI-only trial user, upload path) + follow-up
    state: dict = {}

    def run_sales():
        res = upload_chat(trial, "รวมยอดขายทุก sheet ให้หน่อย แยกตามเดือน และยอดรวมทั้งหมด",
                          [("sales.xlsx", sales_xlsx(), XLSX)])
        state["sales"] = res
        ok = res["source"] == "code" and str(SALES_TOTAL) in digits(res["answer"]) \
            and all(str(v) in digits(res["answer"]) for v in (41650, 33400, 22600))
        record("excel_multisheet_total", ok, res)
    guarded("excel_multisheet_total", run_sales)

    def run_follow_up():
        sid = state["sales"]["session_id"]
        res = stream_chat(trial, "แล้วเดือนไหนยอดขายสูงสุด และสูงกว่าเดือนที่ต่ำสุดกี่บาท", sid)
        ok = ("ม.ค" in res["answer"] or "มกราคม" in res["answer"]) and "19050" in digits(res["answer"])
        record("excel_follow_up_same_session", ok, res)
    guarded("excel_follow_up_same_session", run_follow_up)

    # 4) Pareto chart + Excel output, downloads, owner isolation, history
    def run_pareto():
        res = upload_chat(
            trial,
            "ทำ Pareto chart จากข้อมูลของเสียนี้ พร้อมไฟล์ Excel สรุปจำนวน % และ % สะสม "
            "และบอกว่าปัญหากลุ่มไหนรวมกันประมาณ 80%",
            [("defects.csv", defects_csv(), "text/csv")],
        )
        gen = [a for a in res["attachments"] if a.get("generated")]
        pngs = [a for a in gen if a["content_type"] == "image/png"]
        xlsxs = [a for a in gen if a["filename"].lower().endswith(".xlsx")]
        ok = res["source"] == "code" and bool(pngs) and bool(xlsxs)
        record("pareto_chart_and_excel", ok, res)
        if not ok:
            return

        png = download(trial, pngs[0]["id"])
        size = Image.open(io.BytesIO(png.content)).size if png.status_code == 200 else None
        record("download_png_owner", png.status_code == 200 and png.content[:8] == b"\x89PNG\r\n\x1a\n",
               status=png.status_code, bytes=len(png.content), size=size)

        xl = download(trial, xlsxs[0]["id"])
        nums = []
        if xl.status_code == 200:
            wb = openpyxl.load_workbook(io.BytesIO(xl.content), data_only=True)
            nums = [c for ws in wb for row in ws.iter_rows(values_only=True) for c in row
                    if isinstance(c, (int, float))]
        has_cum = any(abs(n - 87.14) < 0.1 or abs(n - 0.8714) < 0.001 for n in nums)
        record("download_xlsx_owner", xl.status_code == 200 and 120 in nums and has_cum,
               status=xl.status_code, n_numbers=len(nums),
               disposition=xl.headers.get("content-disposition"))

        for viewer in ("e2e_other",):
            r = download(tok[viewer], pngs[0]["id"])
            r2 = download(tok[viewer], xlsxs[0]["id"])
            record(f"download_blocked_for_{viewer}", r.status_code == 403 and r2.status_code == 403,
                   status=[r.status_code, r2.status_code])
        r = httpx.get(f"{BASE}/chat/sessions/{res['session_id']}", headers=hdr(other), timeout=30)
        record("session_hidden_from_other_user", r.status_code == 404, status=r.status_code)

        r = httpx.get(f"{BASE}/chat/sessions/{res['session_id']}", headers=hdr(trial), timeout=30)
        atts = [a for m in r.json()["messages"] for a in m.get("attachments", [])]
        flags = sorted((a["filename"], a["generated"]) for a in atts)
        record("history_keeps_files_with_flags",
               ("defects.csv", False) in flags and sum(1 for _, g in flags if g) == len(gen), files=flags)
    guarded("pareto_chart_and_excel", run_pareto)

    # 5) Concurrency — 3 sandbox jobs at once; another user's chat must stay fast
    def run_concurrency():
        base = stream_chat(other, "สวัสดีครับ")
        with ThreadPoolExecutor(4) as ex:
            jobs = [
                ex.submit(upload_chat, trial, "รวมยอดขายทุก sheet", [("sales.xlsx", sales_xlsx(), XLSX)]),
                ex.submit(upload_chat, admin, "ทำ Pareto chart จากไฟล์นี้", [("defects.csv", defects_csv(), "text/csv")]),
                ex.submit(stream_chat, admin, CALCS[1][1]),
            ]
            time.sleep(4)
            probes = []
            while not all(j.done() for j in jobs) and len(probes) < 4:
                probes.append(stream_chat(other, "ขอบคุณครับ"))
                time.sleep(2)
            done = [j.result() for j in jobs]
        worst = max(p["total_s"] for p in probes) if probes else None
        record("concurrency_other_user_not_blocked",
               bool(probes) and worst < max(15, base["total_s"] * 4) and all(d["source"] == "code" for d in done),
               baseline_s=base["total_s"], probe_total_s=[p["total_s"] for p in probes],
               probe_first_s=[p["first_s"] for p in probes],
               ci_jobs_s=[d["total_s"] for d in done], ci_sources=[d["source"] for d in done])
    guarded("concurrency_other_user_not_blocked", run_concurrency)

    # 6) Regression — features that must keep working
    def run_reg_stream():
        res = stream_chat(other, "อธิบายความแตกต่างระหว่าง CMYK กับ RGB สั้นๆ")
        record("regression_streaming", bool(res["answer"]) and res["source"] not in (None, "code"), res)
    guarded("regression_streaming", run_reg_stream)

    def run_reg_upload_txt():
        res = upload_chat(other, "สรุปไฟล์นี้ให้หน่อย",
                          [("note.txt", "ประชุมวันจันทร์ 10 โมง เรื่องงบประมาณโรงพิมพ์".encode(), "text/plain")])
        record("regression_file_upload_txt", res["source"] == "files", res)
    guarded("regression_file_upload_txt", run_reg_upload_txt)

    def run_reg_xlsx_not_trial():
        res = upload_chat(other, "สรุปไฟล์นี้", [("sales.xlsx", sales_xlsx(), XLSX)])
        record("regression_excel_upload_user_without_flag", res["source"] == "files", res)
    guarded("regression_excel_upload_user_without_flag", run_reg_xlsx_not_trial)

    def run_reg_image():
        res = stream_chat(other, "วาดรูปแมวสีส้มนั่งบนโต๊ะทำงาน สไตล์การ์ตูน")
        record("regression_image_generation", res["source"] == "image_gen" and len(res["attachments"]) == 1, res)
    guarded("regression_image_generation", run_reg_image)

    def run_reg_agent():
        res = stream_chat(admin, "ราคาทองคำแท่งในไทยวันนี้เท่าไหร่")
        record("regression_agent_v1_web", res["source"] == "agent" and "agent_searching" in res["statuses"], res)
    guarded("regression_agent_v1_web", run_reg_agent)

    def run_reg_proofread():
        res = upload_chat(admin, "ช่วยตรวจคำผิดในรูปนี้ให้หน่อย", [("poster.png", typo_png(), "image/png")])
        record("regression_proofreading_v11", res["source"] == "files" and "รางวัน" in res["answer"], res)
    guarded("regression_proofreading_v11", run_reg_proofread)


# ── Phase "off" (backend restarted with both flags off) ─────────────────────

def phase_off(tok: dict[str, str]) -> None:
    admin = tok["e2e_admin"]

    def run_calc_off():
        res = stream_chat(admin, CALCS[0][1])
        record("flag_off_calc_uses_legacy", res["source"] not in ("code", "agent") and not res["statuses"], res)
    guarded("flag_off_calc_uses_legacy", run_calc_off)

    def run_xlsx_off():
        res = upload_chat(admin, "รวมยอดขายทุก sheet", [("sales.xlsx", sales_xlsx(), XLSX)])
        record("flag_off_excel_uses_files_path",
               res["source"] == "files" and not any(a.get("generated") for a in res["attachments"]), res)
    guarded("flag_off_excel_uses_files_path", run_xlsx_off)


def main() -> int:
    global BASE
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--phase", choices=("on", "off"), default="on")
    ap.add_argument("--setup", action="store_true")
    ap.add_argument("--only", nargs="*", default=[], help="run only cases containing these names")
    args = ap.parse_args()
    BASE = args.base.rstrip("/")
    ONLY.extend(args.only)

    if args.setup:
        for u in USERS:
            r = httpx.post(f"{BASE}/auth/register", data={"username": u, "password": PASSWORD}, timeout=30)
            print(u, r.status_code, r.text[:80])
        return 0

    tok = {u: login(u) for u in USERS}
    (phase_on if args.phase == "on" else phase_off)(tok)
    failed = [r["case"] for r in results if not r["ok"]]
    print(f"\n{len(results) - len(failed)}/{len(results)} PASS" + (f" — FAILED: {failed}" if failed else ""))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())

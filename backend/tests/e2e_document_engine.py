"""
End-to-end check of the V3 Document Engine against a RUNNING backend (real
OpenAI; about ฿55–80 per full run — ask before running). Not collected by pytest.

Backend: the real image (LibreOffice for rendering) started with
    DOCUMENT_ENGINE_MODE=admin DOCUMENT_ENGINE_USERS=e2e_trial
    CODE_INTERPRETER_MODE=admin CODE_INTERPRETER_USERS=e2e_trial AGENT_V1_MODE=admin
and the users of tests/e2e_code_interpreter.py (e2e_admin, e2e_trial, e2e_other).

    python tests/e2e_document_engine.py --base http://host:8000              # flags on
    python tests/e2e_document_engine.py --base http://host:8000 --phase off  # after restarting, flags off
    --only <name …>   run just some cases
Downloaded files are kept in --out for review. Prints one JSON line per case.
"""

from __future__ import annotations

import argparse
import io
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import sirigpt_docs as sd  # noqa: E402
from tests import e2e_code_interpreter as e2e  # noqa: E402 — login/stream/upload/record helpers

XLSX = e2e.XLSX
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
OUT = Path("/tmp/v3_outputs")
SALES = {"ม.ค.": 41650, "ก.พ.": 33400, "มี.ค.": 22600}


def save(tok: str, res: dict, ext: str) -> Path | None:
    """Download the first generated file with this extension (owner-checked route)."""
    att = next((a for a in res.get("attachments", []) if a.get("generated")
                and a["filename"].lower().endswith(ext)), None)
    if not att:
        return None
    r = e2e.download(tok, att["id"])
    if r.status_code != 200:
        return None
    path = OUT / f"{int(time.time() * 1000)}_{att['filename']}"
    path.write_bytes(r.content)
    return path


def has_preview(res: dict) -> bool:
    return any(a["filename"].endswith("_ตัวอย่าง.png") for a in res.get("attachments", []))


def docx_text(path: Path) -> str:
    from docx import Document

    d = Document(str(path))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in d.paragraphs] + cells)


# ── inputs made locally (no AI) ─────────────────────────────────────────────

def make_docx() -> bytes:
    doc = sd.new_docx()
    sd.heading(doc, "รายงานประจำเดือน", 1)
    sd.para(doc, "สรุปผลงานฝ่ายผลิตประจำเดือนมีนาคม")
    sd.table(doc, [["งาน", "จำนวน"], ["กล่องลูกฟูก", 120], ["ฉลาก", 450]])
    from PIL import Image

    png = io.BytesIO()
    Image.new("RGB", (400, 200), (91, 44, 159)).save(png, "PNG")
    png.seek(0)
    doc.add_picture(png)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "มี.ค."
    ws.append(["สินค้า", "จำนวน", "ราคา", "ยอดขาย"])
    for i, (n, q, p) in enumerate([("กล่องลูกฟูก", 9, 1200), ("ฉลาก", 16, 450)], start=2):
        ws.append([n, q, p, f"=B{i}*C{i}"])
    ws.append(["รวม", "=SUM(B2:B3)", None, "=SUM(D2:D3)"])
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=3), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
    ws.add_chart(ch, "F2")
    sd.prepare_workbook(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    prs = sd.new_pptx()
    sd.title_slide(prs, "แผนงานไตรมาส 2", "ฝ่ายขาย")
    sd.bullet_slide(prs, "เป้าหมาย", ["เพิ่มยอดขาย 10%", "ลดของเสีย 5%"])
    sd.table_slide(prs, "ตัวเลขสำคัญ", [["รายการ", "ค่า"], ["ยอดขาย", 97650], ["ของเสีย", 350]])
    sd.chart_slide(prs, "แนวโน้ม", list(SALES), {"ยอดขาย": list(SALES.values())})
    sd.bullet_slide(prs, "ขั้นตอนถัดไป", ["ประชุมทีม", "ติดตามผลรายเดือน"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    """2 text pages (one with a ruled table) + 1 scanned Thai page."""
    import fitz
    from PIL import Image, ImageDraw, ImageFont

    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((72, 80), "Sales report Q1/2026 - page 1", fontsize=14)
    p = doc.new_page()
    for r, row in enumerate([["Month", "Sales"], ["Jan", "41650"], ["Feb", "33400"], ["Mar", "22600"]]):
        for c, v in enumerate(row):
            rect = fitz.Rect(72 + c * 150, 100 + r * 24, 72 + (c + 1) * 150, 124 + r * 24)
            p.draw_rect(rect, color=(0, 0, 0), width=1)
            p.insert_text((rect.x0 + 6, rect.y1 - 7), v, fontsize=11)
    img = Image.new("RGB", (1600, 400), "white")
    ImageDraw.Draw(img).text((60, 120), "หมายเหตุ: ยอดขายรวมไตรมาสแรกเพิ่มขึ้น",
                             font=ImageFont.truetype(str(Path(__file__).resolve().parent.parent / "fonts" /
                                                         "Sarabun-Regular.ttf"), 64), fill="black")
    b = io.BytesIO()
    img.save(b, "PNG")
    p = doc.new_page()
    p.insert_image(p.rect, stream=b.getvalue())
    return doc.tobytes()


# ── phase "on" ──────────────────────────────────────────────────────────────

def phase_on(tok: dict[str, str]) -> None:
    trial, other, admin = tok["e2e_trial"], tok["e2e_other"], tok["e2e_admin"]
    state: dict = {}

    def word_report():
        res = e2e.stream_chat(trial, "ทำรายงาน Word ภาษาไทย เรื่องผลการดำเนินงานฝ่ายขายไตรมาส 1/2569 มีบทสรุป "
                                     "ตารางยอดขายรายเดือน (ม.ค. 41,650 บาท ก.พ. 33,400 บาท มี.ค. 22,600 บาท) "
                                     "กราฟแท่งยอดขาย และข้อเสนอแนะ 3 ข้อ")
        path = save(trial, res, ".docx")
        state["word"] = res
        chk = sd.check(str(path)) if path else {"ok": False, "stats": {}}
        s = chk["stats"]
        text = docx_text(path) if path else ""
        ok = (res["source"] == "document" and chk["ok"] and s.get("tables", 0) >= 1 and s.get("images", 0) >= 1
              and s.get("page_numbers") and s.get("thai") and "41,650" in text and has_preview(res))
        e2e.record("word_thai_report_table_image", ok, res, check=s, warnings=chk.get("warnings"),
                   file=path.name if path else None)
    e2e.guarded("word_thai_report", word_report)

    def excel_formulas():
        res = e2e.stream_chat(trial, "สร้างไฟล์ Excel 4 ชีต: ชีต ม.ค. ก.พ. มี.ค. แต่ละชีตมีคอลัมน์ สินค้า จำนวน ราคาต่อหน่วย "
                                     "ยอดขาย(เป็นสูตร) ข้อมูล ม.ค.: กล่อง 10 ชิ้น 1,200 บาท, ฉลาก 25 ชิ้น 450 บาท; "
                                     "ก.พ.: กล่อง 7 ชิ้น 1,200, ฉลาก 30 ชิ้น 450; มี.ค.: กล่อง 9 ชิ้น 1,200, ฉลาก 16 ชิ้น 450 "
                                     "และชีตสรุปที่ดึงยอดแต่ละเดือนด้วยสูตร รวมทั้งไตรมาส พร้อมกราฟแท่ง")
        path = save(trial, res, ".xlsx")
        rep = sd.formula_report(str(path)) if path else {"errors": ["missing"], "values": {}}
        chk = sd.check(str(path)) if path else {"ok": False, "stats": {}}
        expected_total = 10 * 1200 + 25 * 450 + 7 * 1200 + 30 * 450 + 9 * 1200 + 16 * 450   # 63,150
        ok = (res["source"] == "document" and chk["ok"] and not rep["errors"]
              and float(expected_total) in [float(v) for v in rep["values"].values() if isinstance(v, (int, float))]
              and chk["stats"].get("charts", 0) >= 1 and len(chk["stats"].get("sheets", [])) >= 4
              and rep["formulas"] >= 8)
        e2e.record("excel_multisheet_formulas_chart", ok, res, formulas=rep.get("formulas"),
                   verified=rep.get("verified"), errors=rep["errors"][:3], check=chk.get("stats"))
    e2e.guarded("excel_formulas", excel_formulas)

    def pptx_five():
        res = e2e.stream_chat(trial, "ทำ PowerPoint 5 สไลด์ สรุปผลไตรมาส 1: หน้าปก, ประเด็นสำคัญ, ตารางยอดขายรายเดือน "
                                     "(ม.ค. 41,650 ก.พ. 33,400 มี.ค. 22,600), กราฟแนวโน้ม, และแผนไตรมาส 2")
        path = save(trial, res, ".pptx")
        chk = sd.check(str(path)) if path else {"ok": False, "stats": {}, "warnings": []}
        s = chk["stats"]
        ok = (res["source"] == "document" and chk["ok"] and s.get("slides") == 5 and s.get("text_boxes", 0) >= 5
              and (s.get("charts", 0) + s.get("tables", 0)) >= 1 and has_preview(res)
              and not any("ล้น" in w for w in chk["warnings"]))
        e2e.record("pptx_five_editable_slides", ok, res, check=s, warnings=chk.get("warnings"))
    e2e.guarded("pptx_five", pptx_five)

    def pdf_create():
        res = e2e.stream_chat(trial, "ทำเอกสาร PDF ภาษาไทยอย่างน้อย 3 หน้า เรื่องคู่มือความปลอดภัยในโรงพิมพ์ "
                                     "แบ่งเป็น 3 บท บทละหน้า มีหัวข้อและรายการข้อปฏิบัติ")
        path = save(trial, res, ".pdf")
        chk = sd.check(str(path)) if path else {"ok": False, "stats": {}}
        ok = (res["source"] == "document" and chk["ok"] and chk["stats"].get("pages", 0) >= 3
              and chk["stats"].get("thai"))
        e2e.record("pdf_create_multipage_thai", ok, res, check=chk.get("stats"))
    e2e.guarded("pdf_create", pdf_create)

    def pdf_read():
        res = e2e.upload_chat(trial, "อ่าน PDF นี้ทั้งหมด (รวมหน้าที่เป็นภาพสแกน) แล้วทำเป็นไฟล์ Word สรุป พร้อมตารางยอดขาย",
                              [("ยอดขายQ1.pdf", make_pdf(), "application/pdf")])
        path = save(trial, res, ".docx")
        text = docx_text(path).replace(",", "") if path else ""
        ocr = "ยอดขายรวม" in (text + res["answer"]).replace(" ", "")   # only on the scanned page 3
        ok = res["source"] == "document" and all(n in text for n in ("41650", "33400", "22600")) and ocr
        e2e.record("pdf_read_text_table_ocr_to_word", ok, res, found_ocr=ocr)
    e2e.guarded("pdf_read", pdf_read)

    def edit_word():
        orig = OUT / "orig_รายงาน.docx"
        orig.write_bytes(make_docx())
        res = e2e.upload_chat(trial, "แก้ไขไฟล์นี้: เปลี่ยนหัวข้อเป็น 'รายงานประจำเดือนมีนาคม 2569' และเพิ่มย่อหน้าสรุปท้ายเอกสาร",
                              [("รายงาน.docx", orig.read_bytes(), DOCX)])
        path = save(trial, res, ".docx")
        chk = sd.check(str(path), original=str(orig)) if path else {"ok": False, "warnings": ["missing"]}
        lost = [w for w in chk["warnings"] if "ลดลง" in w or "หาย" in w]
        ok = res["source"] == "document" and chk["ok"] and not lost and path is not None \
            and "มีนาคม 2569" in docx_text(path)
        e2e.record("edit_word_keeps_table_and_image", ok, res, lost=lost, file=path.name if path else None)
    e2e.guarded("edit_word", edit_word)

    def edit_excel():
        orig = OUT / "orig_ยอดขาย.xlsx"
        orig.write_bytes(make_xlsx())
        res = e2e.upload_chat(trial, "แก้ไฟล์ Excel นี้: เปลี่ยนจำนวนกล่องลูกฟูกเป็น 12 ชิ้น สูตรอื่นคงเดิม",
                              [("ยอดขาย.xlsx", orig.read_bytes(), XLSX)])
        path = save(trial, res, ".xlsx")
        chk = sd.check(str(path), original=str(orig)) if path else {"ok": False, "warnings": ["missing"]}
        rep = sd.formula_report(str(path)) if path else {"values": {}}
        lost = [w for w in chk["warnings"] if "ลดลง" in w or "หาย" in w]
        ok = res["source"] == "document" and chk["ok"] and not lost \
            and rep["values"].get("มี.ค.!D4") == 12 * 1200 + 16 * 450
        e2e.record("edit_excel_keeps_formulas_and_chart", ok, res, lost=lost, total=rep["values"].get("มี.ค.!D4"))
    e2e.guarded("edit_excel", edit_excel)

    def edit_pptx():
        orig = OUT / "orig_แผนงาน.pptx"
        orig.write_bytes(make_pptx())
        res = e2e.upload_chat(trial, "แก้สไลด์ที่ 2 ของไฟล์นี้: เพิ่มข้อ 'ขยายตลาดต่างจังหวัด' สไลด์อื่นคงเดิม",
                              [("แผนงาน.pptx", orig.read_bytes(), PPTX)])
        path = save(trial, res, ".pptx")
        chk = sd.check(str(path), original=str(orig)) if path else {"ok": False, "warnings": ["missing"], "stats": {}}
        lost = [w for w in chk["warnings"] if "ลดลง" in w or "หาย" in w]
        ok = res["source"] == "document" and chk["ok"] and not lost and chk["stats"].get("slides") == 5
        e2e.record("edit_pptx_keeps_slides_chart_table", ok, res, lost=lost, check=chk.get("stats"))
    e2e.guarded("edit_pptx", edit_pptx)

    def downloads():
        res = state.get("word") or {"attachments": []}   # needs the word_thai_report case
        docs = [a for a in res["attachments"] if a.get("generated")]
        own = [e2e.download(trial, a["id"]).status_code for a in docs]
        others = [e2e.download(other, a["id"]).status_code for a in docs]
        e2e.record("download_owner_only", bool(docs) and set(own) == {200} and set(others) == {403},
                   owner=own, other=others)
    e2e.guarded("downloads", downloads)

    def concurrency():
        base = e2e.stream_chat(other, "สวัสดีครับ")
        with ThreadPoolExecutor(4) as ex:
            jobs = [ex.submit(e2e.stream_chat, trial, "ทำ Word 1 หน้า เรื่องประกาศวันหยุดบริษัท"),
                    ex.submit(e2e.stream_chat, admin, "ทำ PowerPoint 3 สไลด์ เรื่องแนะนำแผนกผลิต"),
                    ex.submit(e2e.stream_chat, trial, "ทำ Excel ตารางเวลาเข้างาน 5 คน พร้อมสูตรรวมชั่วโมง")]
            time.sleep(5)
            probes = []
            while not all(j.done() for j in jobs) and len(probes) < 4:
                probes.append(e2e.stream_chat(other, "ขอบคุณครับ"))
                time.sleep(3)
            done = [j.result() for j in jobs]
        worst = max(p["total_s"] for p in probes) if probes else None
        e2e.record("concurrency_documents_do_not_block_chat",
                   bool(probes) and worst < max(15, base["total_s"] * 4)
                   and all(d["source"] == "document" for d in done),
                   probe_s=[p["total_s"] for p in probes], doc_jobs_s=[d["total_s"] for d in done],
                   sources=[d["source"] for d in done])
    e2e.guarded("concurrency", concurrency)

    # regression of what V3 touched (routing, sandbox, files path)
    def regression():
        checks = [
            ("regression_general_chat", lambda: e2e.stream_chat(other, "อธิบายความต่าง CMYK กับ RGB สั้นๆ"),
             lambda r: r["source"] not in ("document", "code") and r["answer"]),
            ("regression_doc_question_uses_old_path",
             lambda: e2e.upload_chat(trial, "สรุปเนื้อหาไฟล์นี้สั้นๆ", [("a.docx", make_docx(), DOCX)]),
             lambda r: r["source"] == "files"),
            ("regression_code_interpreter_calc",
             lambda: e2e.stream_chat(trial, "ฝากเงิน 100,000 บาท ดอกเบี้ย 5% ทบต้น 10 ปี ได้เท่าไหร่ (ทศนิยม 2 ตำแหน่ง)"),
             lambda r: r["source"] == "code" and "162889.46" in e2e.digits(r["answer"])),
            ("regression_proofreading", lambda: e2e.upload_chat(admin, "ช่วยตรวจคำผิดในรูปนี้ให้หน่อย",
                                                                [("p.png", e2e.typo_png(), "image/png")]),
             lambda r: r["source"] == "files" and "รางวัน" in r["answer"]),
            ("regression_image_generation", lambda: e2e.stream_chat(other, "วาดรูปแมวสีส้มนั่งบนโต๊ะ สไตล์การ์ตูน"),
             lambda r: r["source"] == "image_gen"),
            ("regression_agent_v1_web", lambda: e2e.stream_chat(admin, "ราคาทองคำแท่งวันนี้เท่าไหร่"),
             lambda r: r["source"] == "agent"),
        ]
        for name, run, judge in checks:
            try:
                r = run()
                e2e.record(name, bool(judge(r)), r)
            except Exception as ex:  # noqa: BLE001
                e2e.record(name, False, error=str(ex)[:200])
    e2e.guarded("regression", regression)


def phase_off(tok: dict[str, str]) -> None:
    trial = tok["e2e_trial"]

    def off():
        r = e2e.stream_chat(trial, "ทำรายงาน Word เรื่องประกาศวันหยุด")
        e2e.record("flag_off_no_document_engine", r["source"] != "document" and not any(
            a.get("generated") for a in r["attachments"]), r)
        r = e2e.upload_chat(trial, "แก้ไขไฟล์นี้ เปลี่ยนหัวข้อ", [("a.docx", make_docx(), DOCX)])
        e2e.record("flag_off_docx_upload_uses_files_path", r["source"] == "files", r)
    e2e.guarded("flag_off", off)


def main() -> int:
    global OUT
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--phase", choices=("on", "off"), default="on")
    ap.add_argument("--only", nargs="*", default=[])
    ap.add_argument("--out", default=str(OUT))
    args = ap.parse_args()
    e2e.BASE = args.base.rstrip("/")
    e2e.ONLY.extend(args.only)
    OUT = Path(args.out)
    OUT.mkdir(parents=True, exist_ok=True)
    tok = {u: e2e.login(u) for u in e2e.USERS}
    (phase_on if args.phase == "on" else phase_off)(tok)
    failed = [r["case"] for r in e2e.results if not r["ok"]]
    print(f"\n{len(e2e.results) - len(failed)}/{len(e2e.results)} PASS" + (f" — FAILED: {failed}" if failed else ""))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
